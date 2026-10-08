"""Keep measurement comparisons on matching band definitions and weighting."""

import json
import csv
import contextlib
import io
import tempfile
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import numpy as np

from DUC.results_postprocessing import plot_aeroacoustics as a
from DUC.results_postprocessing import correct_lifter4b_acoustics as calibration
from DUC.results_postprocessing import regenerate_aeroacoustic_plots as regeneration


class CalibrationBandTests(unittest.TestCase):
    fs = 32000.0

    def dataset(self, tone_hz):
        time = np.arange(32000) / self.fs
        pressure = np.sqrt(2) * np.cos(2 * np.pi * tone_hz * time)
        return SimpleNamespace(
            time=time,
            pressure_by_observer={3: pressure},
            thickness_by_observer={3: np.zeros_like(pressure)},
            loading_by_observer={3: pressure},
            selection_windows_by_observer=None,
            pressure_scale_pa=1.0,
            sample_rate_hz=self.fs,
        )

    def test_calibration_uses_the_shared_iec_centers(self):
        nominal = np.asarray([50, 63, 100, 160, 10000], dtype=float)
        centers = calibration.nominal_to_exact_centers(nominal)
        np.testing.assert_allclose(centers, 1000 * 10 ** (np.asarray([-13, -12, -10, -8, 10]) / 10))
        np.testing.assert_array_equal(centers, a.exact_third_octave_centers(nominal))

    def test_calibration_rejects_duplicate_exact_bands_and_unordered_labels(self):
        for invalid in (np.asarray([50, 51]), np.asarray([100, 63]), np.asarray([])):
            with self.assertRaises(ValueError):
                calibration.nominal_to_exact_centers(invalid)

    def test_complete_ten_khz_band_includes_power_above_ten_khz(self):
        dataset = self.dataset(10800)
        actual = calibration.cfd_band_levels(dataset, 3, np.asarray([10000.0]))[0]
        expected = 20 * np.log10(1 / a.ACOUSTIC_PRESSURE_REFERENCE_PA) + a._a_weighting_db(np.asarray([10800]))[0]
        self.assertAlmostEqual(actual, expected, places=5)
        self.assertLess(a.oaspl_a_weighted_db(dataset.pressure_by_observer[3], self.fs), actual - 80)

    def test_band_sums_match_the_measurement_matched_overall(self):
        dataset = self.dataset(1000)
        nominal = np.asarray([50, 63, 80, 100, 125, 160, 200, 250, 315, 400,
                              500, 630, 800, 1000, 1250, 1600, 2000, 2500, 3150,
                              4000, 5000, 6300, 8000, 10000], dtype=float)
        levels = calibration.cfd_band_levels(dataset, 3, calibration.nominal_to_exact_centers(nominal))
        expected = a.measurement_matched_oaspl_db(dataset.pressure_by_observer[3], self.fs, nominal)
        self.assertAlmostEqual(calibration.energy_sum_db(levels), expected, places=12)

    def test_unresolved_or_zero_bands_cannot_define_a_calibration_ratio(self):
        dataset = self.dataset(1000)
        with patch.object(a, 'third_octave_spectrum_at_centers_db', return_value=np.asarray([np.nan])):
            with self.assertRaisesRegex(ValueError, 'unavailable or zero CFD power'):
                calibration.cfd_band_levels(dataset, 3, np.asarray([63.0]))
        with patch.object(a, 'third_octave_spectrum_at_centers_db', return_value=np.asarray([-np.inf])):
            with self.assertRaisesRegex(ValueError, 'unavailable or zero CFD power'):
                calibration.cfd_band_levels(dataset, 3, np.asarray([63.0]))

    def test_silent_bands_are_zero_power_but_missing_bands_are_not_summed(self):
        self.assertEqual(calibration.energy_sum_db(np.asarray([-np.inf])), -np.inf)
        self.assertEqual(calibration.energy_sum_db(np.asarray([0.0, -np.inf])), 0.0)
        for invalid in (np.asarray([np.nan]), np.asarray([np.inf]), np.asarray([])):
            with self.assertRaises(ValueError):
                calibration.energy_sum_db(invalid)

    def test_calibration_rejects_unweighted_measurements_before_loading_cfd(self):
        measurement = SimpleNamespace(spectra_by_observer={
            observer: a.MeasurementSpectrum(np.asarray([1000.0]), np.asarray([70.0]), weighting='Z')
            for observer in calibration.OBSERVERS
        })
        with patch.object(a, 'load_poc2x_1000rpm_measurement_data', return_value=measurement), \
                patch.object(calibration, 'load_cfd') as load:
            with self.assertRaisesRegex(ValueError, 'needs A-weighted measurement bands'):
                calibration.run(Path('measurement.xlsx'), Path('poc2.csv'), Path('lifter.csv'), Path('output'))
        load.assert_not_called()


class RegenerationComparisonTests(unittest.TestCase):
    def test_each_source_resolves_its_own_reference_values(self):
        with patch.object(a, '_timing_for_file', return_value=(2186, 1756)) as timing, \
                patch.object(a, '_reference_values_for_file', return_value=(1.18, 242.65, 298.15, .001)) as references, \
                patch.object(a, 'load_acoustic_dataset') as load:
            csv_file = Path('Pusher/aeroacoustics/step3/total_acoustics_v3.csv')
            regeneration.load_source(csv_file, (2, 3, 4))
        timing.assert_called_once_with(csv_file, None, None)
        references.assert_called_once_with(csv_file, None, None, None, None)
        self.assertEqual(load.call_args.kwargs['rpm'], 2186)
        self.assertEqual(load.call_args.kwargs['reference_temperature_k'], 298.15)
        self.assertEqual(load.call_args.args[1], (2, 3, 4))

    def test_refresh_never_puts_a_weighted_measurements_on_unweighted_spectra(self):
        measurement = SimpleNamespace(spectra_by_observer={3: a.MeasurementSpectrum(
            np.asarray([1000.0]), np.asarray([70.0])
        )})
        dataset = SimpleNamespace(source_file=Path('POC2/total_acoustics_v3.csv'))
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / 'step4_measurements_comparison'
            summary = {'jobs': [{'output_dir': str(output), 'sources': [str(dataset.source_file)],
                                 'revolutions': 'zero:1-3'}]}
            (root / 'aeroacoustic_regeneration_summary.json').write_text(json.dumps(summary), encoding='utf-8')
            with patch.object(regeneration, 'ROOT', root), \
                    patch.object(a, 'load_poc2x_1000rpm_measurement_data', return_value=measurement), \
                    patch.object(regeneration, 'load_source', return_value=dataset), \
                    patch.object(a, 'prepare_acoustic_dataset', return_value=dataset) as prepare, \
                    patch.object(a, 'render_spectrum_plots', return_value=[]), \
                    patch.object(a, 'render_psd_plots', return_value=[]), \
                    patch.object(a, 'render_third_octave_plots', return_value=[]) as render, \
                    patch.object(a, 'render_surface_pair_third_octave_plots', return_value=[]), \
                    patch.object(a, 'APPLY_A_WEIGHTING_TO_SPECTRA', True), \
                    patch.object(a, 'LEGEND_LABEL_OVERRIDE', None), \
                    patch.object(a, 'LEGEND_LABELS_BY_FILE', {}):
                regeneration.refresh_spectral_displays()
        prepare.assert_called_once_with(dataset, 0, 'cfd', a._parse_revolution_range('zero:1-3'))
        self.assertEqual(render.call_count, 2)
        self.assertIs(render.call_args_list[0].kwargs['measurement_spectra_by_observer'],
                      measurement.spectra_by_observer)
        self.assertIsNone(render.call_args_list[1].kwargs['measurement_spectra_by_observer'])

    def test_old_overlay_cleanup_requires_a_matching_replacement(self):
        source = Path('POC2/step4/case-example_total_acoustics_v3.csv')
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            weighted = output / 'a-weighted'
            weighted.mkdir()
            basename = 'case-example_' + a.PLOTS[1].filename
            old = weighted / basename
            flat_old = output / basename
            unrelated = output / 'pressure_time_history_observer_3.png'
            for file in (old, flat_old, unrelated):
                file.touch()
            self.assertEqual(regeneration.remove_obsolete_measurement_oaspl(output, [source]), [])
            replacement = weighted / 'case-example_oaspl_a_weighted_measurement_matched_observers_1_5.png'
            replacement.touch()
            removed = regeneration.remove_obsolete_measurement_oaspl(output, [source])
            self.assertEqual(set(removed), {old, flat_old})
            self.assertFalse(old.exists())
            self.assertFalse(flat_old.exists())
            self.assertTrue(unrelated.exists())
            self.assertTrue(replacement.exists())


class MeasurementRoutingTests(unittest.TestCase):
    nominal = np.asarray([50, 63, 80, 100, 125, 160, 200, 250, 315, 400,
                          500, 630, 800, 1000, 1250, 1600, 2000, 2500, 3150,
                          4000, 5000, 6300, 8000, 10000], dtype=float)

    def measurement(self):
        spectrum = a.MeasurementSpectrum(self.nominal, np.full(24, 70.0))
        return a.Poc2xMeasurementData({3: spectrum}, {3: calibration.energy_sum_db(spectrum.spl_db)})

    def dataset(self):
        fs = 32000.0
        sound_speed = a._speed_of_sound_m_s(298.15)
        scale = 1.1839124829 * sound_speed ** 2
        seconds = np.arange(32000) / fs
        pressure = np.sqrt(2) * (np.cos(2 * np.pi * 1000 * seconds)
                                + np.cos(2 * np.pi * 10800 * seconds)) / scale
        return a.AcousticDataset(
            source_file=Path('POC2/aeroacoustics/step4/case-example_total_acoustics_v3.csv'),
            time=seconds * sound_speed / .001,
            pressure_by_observer={observer: pressure.copy() for observer in (1, 3)},
            thickness_by_observer={observer: pressure * .4 for observer in (1, 3)},
            loading_by_observer={observer: pressure * .6 for observer in (1, 3)},
            rpm=1000, time_steps_per_revolution=8000,
            reference_density_kg_m3=1.1839124829, reference_velocity_m_s=146.0,
            reference_temperature_k=298.15, l_grid_unit_m=.001,
        )

    def patch_renderers(self, stack):
        for name in ('render_spectrum_plots', 'render_psd_plots', 'render_harmonic_content_plots',
                     'render_third_octave_plots', 'render_pressure_time_history_plots',
                     'render_surface_pressure_time_history_plots', 'render_surface_pair_third_octave_plots',
                     'render_average_pressure_time_history_plots', 'render_spectrogram_plots',
                     'render_bpf_component_plots'):
            stack.enter_context(patch.object(a, name, return_value=[]))
        stack.enter_context(patch.object(a, 'render_surface_pair_oaspl_plots', return_value=None))

    def test_single_case_bar_errors_and_csv_use_complete_matched_bands(self):
        dataset, measurement = self.dataset(), self.measurement()
        expected = a.measurement_matched_oaspl_db(a._observer_pressure(dataset, 3), dataset.sample_rate_hz, self.nominal)
        literal = a.oaspl_a_weighted_db(a._observer_pressure(dataset, 3), dataset.sample_rate_hz)
        self.assertGreater(expected - literal, 1.0)
        with tempfile.TemporaryDirectory() as directory, contextlib.ExitStack() as stack:
            output = Path(directory)
            self.patch_renderers(stack)
            stack.enter_context(patch.object(a, 'load_acoustic_dataset', return_value=dataset))
            bars = stack.enter_context(patch.object(a, 'render_bar_plot', return_value=output / 'bar.png'))
            differences = stack.enter_context(patch.object(a, 'print_oaspl_differences'))
            printed = stack.enter_context(patch.object(a, 'print_oaspl_values'))
            stack.enter_context(contextlib.redirect_stdout(io.StringIO()))
            a.make_plots(dataset.source_file, (1, 3), output_dir=output, measurement_data=measurement)
            with (output / 'poc2_test_results_measurement_matched.csv').open(encoding='utf-8') as file:
                rows = list(csv.DictReader(file))
        self.assertIsNone(bars.call_args_list[0].kwargs['measurement_oaspl_dba_by_observer'])
        weighted = bars.call_args_list[1]
        self.assertEqual(weighted.args[1].name, a.MEASUREMENT_MATCHED_OASPL_NAME)
        self.assertAlmostEqual(weighted.args[0][3], expected)
        self.assertIn(1, weighted.args[0])  # All CFD observers get the same comparison bandwidth.
        self.assertAlmostEqual(differences.call_args.args[1][3], expected)
        self.assertAlmostEqual(printed.call_args.args[1]['oaspl_a_weighted'][3], literal)
        self.assertAlmostEqual(float(rows[0]['simulation_dba']), expected, places=2)
        self.assertIn('complete 50–10000 Hz', rows[0]['integration_basis'])

    def test_comparison_routing_retains_literal_values_and_adds_matching_values(self):
        dataset, measurement = self.dataset(), self.measurement()
        expected = a.measurement_matched_oaspl_db(a._observer_pressure(dataset, 3), dataset.sample_rate_hz, self.nominal)
        with tempfile.TemporaryDirectory() as directory, contextlib.ExitStack() as stack:
            output = Path(directory)
            self.patch_renderers(stack)
            stack.enter_context(patch.object(a, 'load_acoustic_dataset', return_value=dataset))
            comparison = stack.enter_context(patch.object(a, 'render_comparison_plot', return_value=output / 'comparison.png'))
            stack.enter_context(contextlib.redirect_stdout(io.StringIO()))
            a.make_comparison_plot([dataset.source_file], (1, 3), output_dir=output, measurement_data=measurement)
        values = comparison.call_args.args[0][0][1]
        self.assertIn('oaspl_a_weighted', values)
        self.assertAlmostEqual(values[a.MEASUREMENT_MATCHED_OASPL_NAME][3], expected)
        np.testing.assert_array_equal(comparison.call_args.kwargs['measurement_band_centers_hz'], self.nominal)

    def test_comparison_renderer_uses_matched_values_on_the_weighted_axis(self):
        dataset = self.dataset()
        values = {'oaspl_unweighted': {1: 60, 3: 61}, 'oaspl_a_weighted': {1: 50, 3: 51},
                  a.MEASUREMENT_MATCHED_OASPL_NAME: {1: 90, 3: 91}}
        with tempfile.TemporaryDirectory() as directory, patch.object(a, '_save_figure') as save:
            output = a.render_comparison_plot([(dataset, values)], (1, 3), Path(directory),
                                             {3: 80}, measurement_band_centers_hz=self.nominal)
        np.testing.assert_array_equal(save.call_args.args[0].axes[1].lines[0].get_ydata(), [90, 91])
        self.assertIn('complete 50–10000 Hz', save.call_args.args[0].axes[1].get_ylabel())
        self.assertIn('measurement_matched', output.name)

    def test_direct_renderers_reject_weighting_and_bandwidth_mismatches(self):
        with self.assertRaisesRegex(ValueError, 'measurement-matched'):
            a.render_bar_plot({3: 50}, a.PLOTS[1], Path('POC2/case.csv'),
                              measurement_oaspl_dba_by_observer={3: 60})
        with patch.object(a, 'APPLY_A_WEIGHTING_TO_SPECTRA', False):
            with self.assertRaisesRegex(ValueError, 'weighting/bands do not match'):
                a.render_third_octave_plots([], (3,), measurement_spectra_by_observer=self.measurement().spectra_by_observer)

    def test_excel_overall_must_equal_its_twenty_four_band_energies(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        for observer, column in a.POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER.items():
            sheet.cell(3, column, f'{a.OBSERVER_ANGLES_DEG[observer]} 1000 rpm')
            for row, frequency in enumerate(self.nominal, 4):
                sheet.cell(row, 1, frequency)
                sheet.cell(row, column, 70.0)
            sheet.cell(28, column, calibration.energy_sum_db(np.full(24, 70.0)))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'measurement.xlsx'
            workbook.save(path)
            loaded = a.load_poc2x_1000rpm_measurement_data(path)
            self.assertEqual(set(loaded.spectra_by_observer), {2, 3, 4})
            column = a.POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER[3]
            sheet.cell(28, column, sheet.cell(28, column).value + 1.0)
            workbook.save(path)
            with self.assertRaisesRegex(ValueError, 'does not equal its band-energy sum'):
                a.load_poc2x_1000rpm_measurement_data(path)
        workbook.close()


if __name__ == '__main__':
    unittest.main()
