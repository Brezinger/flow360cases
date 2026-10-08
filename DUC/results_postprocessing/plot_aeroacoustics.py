from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Callable

try:
    from bunch_surface_acoustics import bunch_surface_resolved_acoustics
except ImportError:
    bunch_surface_resolved_acoustics = None

import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter, LogLocator, MultipleLocator
import numpy as np
from scipy.signal import spectrogram, welch


SCRIPT_DIR = Path(__file__).resolve().parent
RESULTS_DIR = SCRIPT_DIR / "POC2"
AEROACOUSTICS_DIR = RESULTS_DIR / "aeroacoustics"
PLOTS_DIR = RESULTS_DIR / "aeroacoustic_plots"

ACOUSTIC_PRESSURE_REFERENCE_PA = 20.0e-6
AIR_SPECIFIC_HEAT_RATIO = 1.4
AIR_GAS_CONSTANT_J_KG_K = 287.05287
DEFAULT_OBSERVERS = (1, 2, 3, 4, 5)
OBSERVER_ANGLES_DEG = {
    1: 0,
    2: 90,
    3: 105,
    4: 135,
    5: 180,
}
FLOW360_OBSERVER_NUMBERING_OFFSET = -1
PUSHER_OBSERVER_ANGLES_DEG = (20, 45, 90, 135, 160)
DEFAULT_RPM = 1063.677
POC2_RPM = 1000.0
Q40LT_RPM = 1002.0
POC2_PROPELLER_RADIUS_M = 1.4
STEP3_TIME_STEPS_PER_REVOLUTION = 120 * 32
STEP4_TIME_STEPS_PER_REVOLUTION = STEP3_TIME_STEPS_PER_REVOLUTION * 25 / 12
DEFAULT_TIME_STEPS_PER_REVOLUTION = STEP4_TIME_STEPS_PER_REVOLUTION
DEFAULT_BLADE_COUNT = 4
LOCAL_RESULTS_REFERENCE_DENSITY_KG_M3 = 1.064099
LOCAL_RESULTS_REFERENCE_VELOCITY_M_S = 155.9432
LOCAL_RESULTS_REFERENCE_TEMPERATURE_K = 303.275
LOCAL_RESULTS_L_GRID_UNIT_M = 0.001
POC2_REFERENCE_VELOCITY_M_S = 2.0 * np.pi * POC2_PROPELLER_RADIUS_M * POC2_RPM / 60.0
POC2X2_REFERENCE_DENSITY_KG_M3 = 1.149
POC2X2_REFERENCE_VELOCITY_M_S = 146.857683
POC2X2_REFERENCE_TEMPERATURE_K = 283.275
POC2X2_L_GRID_UNIT_M = 1.0
Q40LT_REFERENCE_DENSITY_KG_M3 = 1.149
Q40LT_REFERENCE_VELOCITY_M_S = 146.8587100942248
Q40LT_REFERENCE_TEMPERATURE_K = 283.275
Q40LT_L_GRID_UNIT_M = 1.0
CASE_TIMING_BY_PREFIX = {
    "step3": (DEFAULT_RPM, STEP3_TIME_STEPS_PER_REVOLUTION),
    "step4": (DEFAULT_RPM, STEP4_TIME_STEPS_PER_REVOLUTION),
}
COMPARISON_OUTPUT_FILE = "oaspl_observers_1_5_comparison.png"
OASPL_MIN_FREQUENCY_HZ = 40.0
OASPL_MAX_FREQUENCY_HZ = 10_000.0
SPECTRUM_MIN_FREQUENCY_HZ = 40.0
SPECTRUM_MAX_FREQUENCY_HZ = 10_000.0
SPECTRUM_X_TICKS_HZ = (40, 60, 80, 100, 200, 400, 600, 800, 1000, 2000, 4000, 6000, 8000, 10000)
SPECTRUM_Y_LIMITS_DB = (-20.0, 70.0)
SPECTROGRAM_MAX_FREQUENCY_HZ = 5_000.0
SPECTROGRAM_LIMITS_DB = (0.0, 120.0)
SPECTROGRAM_NPERSEG = 512
HARMONIC_INDICES = tuple(range(1, 11))
HARMONIC_Y_LIMITS_DB = (0.0, 110.0)
SURFACE_TIME_HISTORY_GROUPS = ("blade1", "blade2")
SURFACE_THIRD_OCTAVE_GROUP_PAIRS = (
    ("blade1+blade3", ("blade1", "blade3")),
    ("blade2+blade4", ("blade2", "blade4")),
)
POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER = {
    2: 5,
    3: 6,
    4: 7,
}
APPLY_A_WEIGHTING_TO_SPECTRA = True
LEGEND_LABEL_OVERRIDE: str | None = None
LEGEND_LABELS_BY_FILE: dict[Path, str] = {}
MEASUREMENT_LABEL = "Experiment 1000 RPM"
NPERSEG = "All"
WINDOW = "hann"
CASE_ID_PATTERN = re.compile(r"case-[0-9a-fA-F-]+")


def _default_line_colors() -> list[str]:
    return plt.rcParams["axes.prop_cycle"].by_key()["color"]


def _is_local_step4(source_file: Path) -> bool:
    path_parts = {part.lower() for part in source_file.resolve().parts}
    return _step_label(source_file) == "step4" and "poc2x2" not in path_parts


def _line_color(source_file: Path, index: int) -> str:
    del source_file
    colors = _default_line_colors()
    return colors[index % len(colors)]


@dataclass(frozen=True)
class RevolutionRange:
    start_index: int
    end_index: int
    label: str

    @property
    def count(self) -> int:
        return self.end_index - self.start_index


@dataclass(frozen=True)
class AcousticDataset:
    source_file: Path
    time: np.ndarray
    pressure_by_observer: dict[int, np.ndarray]
    thickness_by_observer: dict[int, np.ndarray]
    loading_by_observer: dict[int, np.ndarray]
    rpm: float
    time_steps_per_revolution: float
    reference_density_kg_m3: float
    reference_velocity_m_s: float
    reference_temperature_k: float
    l_grid_unit_m: float
    revolution_range: RevolutionRange | None = None
    source_files: tuple[Path, ...] = ()

    @property
    def sample_spacing_s(self) -> float:
        if len(self.time) < 2:
            raise ValueError("At least two time samples are required to compute the acoustic sample spacing.")
        return float(np.mean(np.diff(self.time * (self.l_grid_unit_m / self.speed_of_sound_m_s))))

    @property
    def sample_rate_hz(self) -> float:
        return 1.0 / self.sample_spacing_s

    @property
    def time_seconds(self) -> np.ndarray:
        return self.time * (self.l_grid_unit_m / self.speed_of_sound_m_s)

    @property
    def pressure_scale_pa(self) -> float:
        return self.reference_density_kg_m3 * self.speed_of_sound_m_s**2

    @property
    def speed_of_sound_m_s(self) -> float:
        return _speed_of_sound_m_s(self.reference_temperature_k)


@dataclass(frozen=True)
class SurfaceResolvedDataset:
    group_name: str
    dataset: AcousticDataset
    case_source_file: Path


@dataclass(frozen=True)
class PlotDefinition:
    name: str
    filename: str
    evaluator: Callable[[np.ndarray, float], float]
    ylabel: str
    title: str


@dataclass(frozen=True)
class MeasurementSpectrum:
    frequencies_hz: np.ndarray
    spl_db: np.ndarray


@dataclass(frozen=True)
class Poc2xMeasurementData:
    spectra_by_observer: dict[int, MeasurementSpectrum]
    oaspl_dba_by_observer: dict[int, float]


def _csv_observer_index(user_observer_index: int) -> int:
    return user_observer_index + FLOW360_OBSERVER_NUMBERING_OFFSET


def _observer_angle_deg(source_file: Path, observer: int) -> int | None:
    if "pusher" in {part.lower() for part in Path(source_file).parts}:
        csv_index = _csv_observer_index(observer)
        if 0 <= csv_index < len(PUSHER_OBSERVER_ANGLES_DEG):
            return PUSHER_OBSERVER_ANGLES_DEG[csv_index]
        return None
    return OBSERVER_ANGLES_DEG.get(observer)


def _default_acoustic_files() -> list[Path]:
    results_dir = AEROACOUSTICS_DIR / "step4"
    candidates = sorted(results_dir.glob("*_total_acoustics_v*.csv"))
    if not candidates:
        raise FileNotFoundError(f"No Flow360 step4 total-acoustics CSV files found in {results_dir}")
    return candidates


def _speed_of_sound_m_s(reference_temperature_k: float) -> float:
    return float(
        np.sqrt(
            AIR_SPECIFIC_HEAT_RATIO
            * AIR_GAS_CONSTANT_J_KG_K
            * reference_temperature_k
        )
    )


def _samples_per_revolution(
    time_values: np.ndarray,
    rpm: float,
    reference_temperature_k: float,
    l_grid_unit_m: float,
) -> int:
    if len(time_values) < 2:
        raise ValueError("At least two time samples are required to select revolutions.")

    speed_of_sound_m_s = _speed_of_sound_m_s(reference_temperature_k)
    sample_spacing_s = float(np.mean(np.diff(time_values * (l_grid_unit_m / speed_of_sound_m_s))))
    samples_per_revolution = int(round((60.0 / rpm) / sample_spacing_s))
    if samples_per_revolution < 1:
        raise ValueError("Could not infer a valid acoustic sample count per revolution.")
    return samples_per_revolution


def _slice_revolutions(
    values: np.ndarray,
    revolution_range: RevolutionRange | None,
    samples_per_revolution: int,
    source_file: Path,
) -> np.ndarray:
    if revolution_range is None:
        return values

    start_sample = revolution_range.start_index * samples_per_revolution
    end_sample = revolution_range.end_index * samples_per_revolution
    if end_sample > len(values):
        available_revolutions = len(values) // samples_per_revolution
        raise ValueError(
            f"Requested revolutions {revolution_range.label} from {source_file}, but only "
            f"{available_revolutions} complete revolutions are available."
        )
    return values[start_sample:end_sample]


def load_acoustic_dataset(
    csv_file: Path | None = None,
    observers: tuple[int, ...] = DEFAULT_OBSERVERS,
    rpm: float = DEFAULT_RPM,
    time_steps_per_revolution: float = DEFAULT_TIME_STEPS_PER_REVOLUTION,
    reference_density_kg_m3: float = LOCAL_RESULTS_REFERENCE_DENSITY_KG_M3,
    reference_velocity_m_s: float = LOCAL_RESULTS_REFERENCE_VELOCITY_M_S,
    reference_temperature_k: float = LOCAL_RESULTS_REFERENCE_TEMPERATURE_K,
    l_grid_unit_m: float = LOCAL_RESULTS_L_GRID_UNIT_M,
    revolution_range: RevolutionRange | None = None,
) -> AcousticDataset:
    csv_file = _default_acoustic_files()[-1] if csv_file is None else Path(csv_file)
    pressure_columns = {
        observer: f"observer_{_csv_observer_index(observer)}_pressure"
        for observer in observers
    }
    thickness_columns = {
        observer: f"observer_{_csv_observer_index(observer)}_thickness"
        for observer in observers
    }
    loading_columns = {
        observer: f"observer_{_csv_observer_index(observer)}_loading"
        for observer in observers
    }
    time_values: list[float] = []
    pressure_values = {observer: [] for observer in observers}
    thickness_values = {observer: [] for observer in observers}
    loading_values = {observer: [] for observer in observers}

    with csv_file.open(newline="") as file:
        reader = csv.DictReader(file, skipinitialspace=True)
        required_columns = [
            *pressure_columns.values(),
            *thickness_columns.values(),
            *loading_columns.values(),
        ]
        missing_columns = [column for column in required_columns if column not in reader.fieldnames]
        if missing_columns:
            raise KeyError(f"Missing acoustic pressure columns in {csv_file}: {missing_columns}")

        for row in reader:
            time_values.append(float(row["time"]))
            for observer, column in pressure_columns.items():
                pressure_values[observer].append(float(row[column]))
            for observer, column in thickness_columns.items():
                thickness_values[observer].append(float(row[column]))
            for observer, column in loading_columns.items():
                loading_values[observer].append(float(row[column]))

    if not time_values:
        raise ValueError(f"No acoustic samples found in {csv_file}")

    time_array = np.asarray(time_values, dtype=float)
    pressure_arrays = {
        observer: np.asarray(values, dtype=float)
        for observer, values in pressure_values.items()
    }
    thickness_arrays = {
        observer: np.asarray(values, dtype=float)
        for observer, values in thickness_values.items()
    }
    loading_arrays = {
        observer: np.asarray(values, dtype=float)
        for observer, values in loading_values.items()
    }

    if revolution_range is not None:
        samples_per_revolution = _samples_per_revolution(
            time_array,
            rpm,
            reference_temperature_k,
            l_grid_unit_m,
        )
        time_array = _slice_revolutions(
            time_array,
            revolution_range,
            samples_per_revolution,
            csv_file,
        )
        pressure_arrays = {
            observer: _slice_revolutions(values, revolution_range, samples_per_revolution, csv_file)
            for observer, values in pressure_arrays.items()
        }
        thickness_arrays = {
            observer: _slice_revolutions(values, revolution_range, samples_per_revolution, csv_file)
            for observer, values in thickness_arrays.items()
        }
        loading_arrays = {
            observer: _slice_revolutions(values, revolution_range, samples_per_revolution, csv_file)
            for observer, values in loading_arrays.items()
        }

    return AcousticDataset(
        source_file=csv_file,
        time=time_array,
        pressure_by_observer=pressure_arrays,
        thickness_by_observer=thickness_arrays,
        loading_by_observer=loading_arrays,
        rpm=rpm,
        time_steps_per_revolution=time_steps_per_revolution,
        reference_density_kg_m3=reference_density_kg_m3,
        reference_velocity_m_s=reference_velocity_m_s,
        reference_temperature_k=reference_temperature_k,
        l_grid_unit_m=l_grid_unit_m,
        revolution_range=revolution_range,
    )


def combine_acoustic_datasets(datasets: list[AcousticDataset]) -> AcousticDataset:
    """Sum continuation/contribution files on their original shared time grid."""
    if not datasets:
        raise ValueError("At least one acoustic dataset is required.")
    first = datasets[0]
    if len(first.time) < 2:
        raise ValueError("At least two timestamps are required to combine acoustic data.")
    spacing = float(np.mean(np.diff(first.time)))
    if spacing <= 0 or not np.isfinite(spacing):
        raise ValueError("Acoustic timestamps must increase.")
    origin = min(float(dataset.time[0]) for dataset in datasets)
    observers = set(first.pressure_by_observer)
    attributes = (
        "rpm", "time_steps_per_revolution", "reference_density_kg_m3",
        "reference_velocity_m_s", "reference_temperature_k", "l_grid_unit_m",
    )
    components = ("pressure_by_observer", "thickness_by_observer", "loading_by_observer")
    grids = []
    for dataset in datasets:
        if any(not np.isclose(getattr(dataset, name), getattr(first, name), rtol=1e-9)
               for name in attributes):
            raise ValueError("Combined files require matching RPM and reference quantities.")
        if any(set(getattr(dataset, component)) != observers for component in components):
            raise ValueError("Combined files require matching observer columns.")
        if len(dataset.time) < 2 or not np.all(np.diff(dataset.time) > 0):
            raise ValueError(f"Non-increasing acoustic timestamps in {dataset.source_file}")
        if not np.isclose(np.mean(np.diff(dataset.time)), spacing, rtol=1e-5, atol=0):
            raise ValueError("Combined files have incompatible acoustic sampling intervals.")
        grid = np.rint((dataset.time - origin) / spacing).astype(int)
        if (
            np.any(np.diff(grid) != 1)
            or np.max(np.abs(dataset.time - (origin + grid * spacing))) > spacing * 0.01
        ):
            raise ValueError(f"Acoustic timestamps do not align to the common grid: {dataset.source_file}")
        grids.append(grid)
    count = max(int(grid[-1]) for grid in grids) + 1
    summed = {component: {observer: np.zeros(count) for observer in observers} for component in components}
    coverage = {observer: np.zeros(count, dtype=bool) for observer in observers}
    for dataset, grid in zip(datasets, grids):
        for observer in observers:
            active = np.flatnonzero(np.any([
                getattr(dataset, component)[observer] != 0 for component in components
            ], axis=0))
            if active.size:
                coverage[observer][grid[active[0]:active[-1] + 1]] = True
            for component in components:
                summed[component][observer][grid] += getattr(dataset, component)[observer]
    for observer, covered in coverage.items():
        active = np.flatnonzero(covered)
        if active.size and not np.all(covered[active[0]:active[-1] + 1]):
            raise ValueError(f"Observer {observer} has a gap between acoustic signal windows.")
    return replace(
        first, time=origin + np.arange(count) * spacing,
        **summed, source_files=tuple(dataset.source_file for dataset in datasets),
    )


def prepare_acoustic_dataset(
    dataset: AcousticDataset,
    n_timesteps_skip: int = 0,
    skip_unit: str = "cfd",
    revolution_range: RevolutionRange | None = None,
) -> AcousticDataset:
    """Cut each observer's initial signal interval, keeping components aligned."""
    if n_timesteps_skip < 0 or skip_unit not in ("cfd", "acoustic"):
        raise ValueError("Skip count must be nonnegative and units must be cfd or acoustic.")
    skip_samples = n_timesteps_skip
    if skip_unit == "cfd":
        cfd_spacing = 60.0 / (dataset.rpm * dataset.time_steps_per_revolution)
        skip_samples = int(round(n_timesteps_skip * cfd_spacing / dataset.sample_spacing_s))
    components = ("pressure_by_observer", "thickness_by_observer", "loading_by_observer")
    prepared = {component: {} for component in components}
    samples_per_revolution = int(round(dataset.sample_rate_hz * 60.0 / dataset.rpm))
    print(f"Skipping {n_timesteps_skip} {skip_unit} steps = {skip_samples} acoustic samples per observer.")
    for observer in dataset.pressure_by_observer:
        active = np.flatnonzero(np.any([
            getattr(dataset, component)[observer] != 0 for component in components
        ], axis=0))
        if not active.size:
            raise ValueError(f"Observer {observer} has no nonzero acoustic signal.")
        start, stop = int(active[0]) + skip_samples, int(active[-1]) + 1
        if revolution_range is not None:
            end = start + revolution_range.end_index * samples_per_revolution
            start += revolution_range.start_index * samples_per_revolution
            if end > stop:
                raise ValueError(f"Observer {observer} has too few retained revolutions.")
            stop = end
        if stop - start < 8:
            raise ValueError(f"Observer {observer} has fewer than eight samples after the cut.")
        for component in components:
            values = getattr(dataset, component)[observer].copy()
            values[:start] = 0
            values[stop:] = 0
            prepared[component][observer] = values
        print(f"  Observer {observer}: {stop - start} retained samples, "
              f"{(stop - start) * dataset.sample_spacing_s:.6f} s")
    return replace(dataset, **prepared, revolution_range=revolution_range)


def load_surface_resolved_datasets(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    groups: tuple[str, ...] = SURFACE_TIME_HISTORY_GROUPS,
    observers: tuple[int, ...] = DEFAULT_OBSERVERS,
) -> list[SurfaceResolvedDataset]:
    surface_datasets = []
    for total_dataset, _ in results_by_case:
        if len(total_dataset.source_files) > 1:
            print("Skipping surface-resolved plots: combined total files have no combined surface dataset.")
            continue
        for group_name in groups:
            surface_file = _surface_file_for_case(total_dataset.source_file, group_name)
            if surface_file is None:
                print(f"Skipping missing surface-resolved file for {total_dataset.source_file.name}: {group_name}")
                continue
            surface_dataset = load_acoustic_dataset(
                surface_file,
                observers=observers,
                rpm=total_dataset.rpm,
                time_steps_per_revolution=total_dataset.time_steps_per_revolution,
                reference_density_kg_m3=total_dataset.reference_density_kg_m3,
                reference_velocity_m_s=total_dataset.reference_velocity_m_s,
                reference_temperature_k=total_dataset.reference_temperature_k,
                l_grid_unit_m=total_dataset.l_grid_unit_m,
                revolution_range=total_dataset.revolution_range,
            )
            surface_datasets.append(
                SurfaceResolvedDataset(
                    group_name=group_name,
                    dataset=surface_dataset,
                    case_source_file=total_dataset.source_file,
                )
            )
    return surface_datasets


def _fluctuating_pressure(pressure: np.ndarray) -> np.ndarray:
    return pressure - np.mean(pressure)


def _filter_zero_signal_edges(pressure: np.ndarray) -> np.ndarray:
    pressure = np.asarray(pressure, dtype=float)
    nonzero_indices = np.where(pressure != 0.0)[0]
    if nonzero_indices.size == 0:
        return pressure
    return pressure[nonzero_indices[0]:nonzero_indices[-1] + 1]


def _pressure_time_history(dataset: AcousticDataset, observer: int) -> tuple[np.ndarray, np.ndarray]:
    """Clip edge padding together with its timestamps; retain interior zeros."""
    pressure_pa = dataset.pressure_by_observer[observer] * dataset.pressure_scale_pa
    nonzero_indices = np.flatnonzero(pressure_pa != 0.0)
    valid = (
        slice(nonzero_indices[0], nonzero_indices[-1] + 1)
        if nonzero_indices.size else slice(None)
    )
    time_seconds = dataset.time_seconds[valid]
    return time_seconds - time_seconds[0], pressure_pa[valid]


def _db_from_power(mean_square_pressure: np.ndarray | float) -> np.ndarray | float:
    mean_square_pressure = np.asarray(mean_square_pressure, dtype=float)
    spl_db = 10.0 * np.log10(
        np.maximum(mean_square_pressure, np.finfo(float).tiny)
        / ACOUSTIC_PRESSURE_REFERENCE_PA**2
    )
    if spl_db.ndim == 0:
        return float(spl_db)
    return spl_db


def _nperseg(sample_count: int) -> int:
    if isinstance(NPERSEG, str) and NPERSEG.strip().lower() == "all":
        return sample_count
    return min(int(NPERSEG), sample_count)


def oaspl_unweighted_db(pressure: np.ndarray, sample_rate_hz: float) -> float:
    frequencies, mean_square_by_bin = _band_limited_mean_square_spectrum(pressure, sample_rate_hz)
    del frequencies
    mean_square_pressure = float(np.sum(mean_square_by_bin))
    return float(_db_from_power(mean_square_pressure))


def _a_weighting_db(frequency_hz: np.ndarray) -> np.ndarray:
    frequency = np.asarray(frequency_hz, dtype=float)
    frequency_squared = np.square(frequency)
    ra = (12200.0**2 * frequency_squared**2) / (
        (frequency_squared + 20.6**2)
        * (frequency_squared + 12200.0**2)
        * np.sqrt((frequency_squared + 107.7**2) * (frequency_squared + 737.9**2))
    )
    ra = np.maximum(ra, 1.0e-12)
    ra_1khz = (12200.0**2 * (1000.0**2) ** 2) / (
        (1000.0**2 + 20.6**2)
        * (1000.0**2 + 12200.0**2)
        * np.sqrt((1000.0**2 + 107.7**2) * (1000.0**2 + 737.9**2))
    )
    weighting = 20.0 * np.log10(ra / ra_1khz)
    return np.where(frequency > 0.0, weighting, -np.inf)


def _weighting_db(frequency_hz: np.ndarray, apply_a_weighting: bool) -> np.ndarray:
    if apply_a_weighting:
        return _a_weighting_db(frequency_hz)
    return np.zeros_like(np.asarray(frequency_hz, dtype=float))


def _welch_mean_square_spectrum(pressure: np.ndarray, sample_rate_hz: float) -> tuple[np.ndarray, np.ndarray]:
    pressure = _filter_zero_signal_edges(pressure)
    sample_count = len(pressure)
    if sample_count < 8:
        raise ValueError("Too few samples for Welch calculation.")

    nperseg = _nperseg(sample_count)
    frequencies, mean_square_by_bin = welch(
        _fluctuating_pressure(pressure),
        fs=sample_rate_hz,
        window=WINDOW,
        nperseg=nperseg,
        noverlap=nperseg // 2,
        scaling="spectrum",
    )
    return frequencies, mean_square_by_bin


def _band_limited_mean_square_spectrum(pressure: np.ndarray, sample_rate_hz: float) -> tuple[np.ndarray, np.ndarray]:
    frequencies, mean_square_by_bin = _welch_mean_square_spectrum(pressure, sample_rate_hz)
    frequency_mask = (
        (frequencies >= OASPL_MIN_FREQUENCY_HZ)
        & (frequencies <= OASPL_MAX_FREQUENCY_HZ)
    )
    if not np.any(frequency_mask):
        raise ValueError("No frequencies remain inside the configured frequency band.")
    return frequencies[frequency_mask], mean_square_by_bin[frequency_mask]


def oaspl_a_weighted_db(pressure: np.ndarray, sample_rate_hz: float) -> float:
    frequencies, mean_square_by_bin = _band_limited_mean_square_spectrum(pressure, sample_rate_hz)
    linear_a_weight = 10.0 ** (_weighting_db(frequencies, apply_a_weighting=True) / 10.0)
    weighted_mean_square_pressure = float(np.sum(mean_square_by_bin * linear_a_weight))
    return float(_db_from_power(weighted_mean_square_pressure))


def spl_spectrum_db(
    pressure: np.ndarray,
    sample_rate_hz: float,
    apply_a_weighting: bool | None = None,
) -> tuple[np.ndarray, np.ndarray]:
    if apply_a_weighting is None:
        apply_a_weighting = APPLY_A_WEIGHTING_TO_SPECTRA
    frequencies, mean_square_by_bin = _band_limited_mean_square_spectrum(pressure, sample_rate_hz)
    linear_weight = 10.0 ** (_weighting_db(frequencies, apply_a_weighting) / 10.0)
    return frequencies, _db_from_power(mean_square_by_bin * linear_weight)


def tone_spl_db(pressure: np.ndarray, sample_rate_hz: float, target_frequency_hz: float) -> tuple[float, float]:
    frequencies, mean_square_by_bin = _welch_mean_square_spectrum(pressure, sample_rate_hz)
    valid_frequency_mask = frequencies > 0.0
    frequencies = frequencies[valid_frequency_mask]
    mean_square_by_bin = mean_square_by_bin[valid_frequency_mask]
    nearest_index = int(np.argmin(np.abs(frequencies - target_frequency_hz)))
    spl_db = _db_from_power(mean_square_by_bin[nearest_index])
    return float(frequencies[nearest_index]), float(spl_db)


def harmonic_third_octave_spl_db(
    pressure: np.ndarray,
    sample_rate_hz: float,
    target_frequency_hz: float,
) -> float:
    frequencies, mean_square_by_bin = _welch_mean_square_spectrum(pressure, sample_rate_hz)
    band_factor = 2.0 ** (1.0 / 6.0)
    band_mask = (
        (frequencies >= target_frequency_hz / band_factor)
        & (frequencies < target_frequency_hz * band_factor)
    )
    return float(_db_from_power(np.sum(mean_square_by_bin[band_mask])))


def _third_octave_centers(fmin_hz: float, fmax_hz: float) -> np.ndarray:
    centers = []
    k_min = int(np.floor(3.0 * np.log2(max(fmin_hz, 1.0e-12) / 1000.0)))
    k_max = int(np.ceil(3.0 * np.log2(fmax_hz / 1000.0)))
    for k in range(k_min, k_max + 1):
        center_frequency = 1000.0 * (2.0 ** (k / 3.0))
        if center_frequency >= fmin_hz:
            centers.append(center_frequency)
    return np.asarray(centers, dtype=float)


def _third_octave_bands(fmin_hz: float, fmax_hz: float) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    centers = _third_octave_centers(fmin_hz, fmax_hz)
    factor = 2.0 ** (1.0 / 6.0)
    return centers, centers / factor, centers * factor


def third_octave_spectrum_db(
    pressure: np.ndarray,
    sample_rate_hz: float,
    apply_a_weighting: bool | None = None,
) -> tuple[np.ndarray, np.ndarray]:
    if apply_a_weighting is None:
        apply_a_weighting = APPLY_A_WEIGHTING_TO_SPECTRA
    frequencies, mean_square_by_bin = _band_limited_mean_square_spectrum(pressure, sample_rate_hz)
    linear_weight = 10.0 ** (_weighting_db(frequencies, apply_a_weighting) / 10.0)
    center_frequencies, lower_band_edges, upper_band_edges = _third_octave_bands(
        OASPL_MIN_FREQUENCY_HZ,
        OASPL_MAX_FREQUENCY_HZ,
    )
    mean_square_by_band = []

    for lower_edge, upper_edge in zip(lower_band_edges, upper_band_edges):
        band_mask = (frequencies >= lower_edge) & (frequencies < upper_edge)
        mean_square_by_band.append(float(np.sum(mean_square_by_bin[band_mask] * linear_weight[band_mask])))

    return center_frequencies, _db_from_power(np.asarray(mean_square_by_band, dtype=float))


PLOTS = (
    PlotDefinition(
        name="oaspl_unweighted",
        filename="oaspl_unweighted_observers_1_5.png",
        evaluator=oaspl_unweighted_db,
        ylabel="OASPL [dB]",
        title="Unweighted OASPL",
    ),
    PlotDefinition(
        name="oaspl_a_weighted",
        filename="oaspl_a_weighted_observers_1_5.png",
        evaluator=oaspl_a_weighted_db,
        ylabel="A-weighted OASPL [dBA]",
        title="A-weighted OASPL",
    ),
)


def evaluate_plot(dataset: AcousticDataset, plot: PlotDefinition) -> dict[int, float]:
    return {
        observer: plot.evaluator(pressure * dataset.pressure_scale_pa, dataset.sample_rate_hz)
        for observer, pressure in dataset.pressure_by_observer.items()
    }


def print_oaspl_values(
    dataset: AcousticDataset,
    values_by_plot: dict[str, dict[int, float]],
    observers: tuple[int, ...],
) -> None:
    print(f"OASPL values for {_case_label(dataset.source_file)}:")
    for observer in observers:
        unweighted = values_by_plot["oaspl_unweighted"][observer]
        a_weighted = values_by_plot["oaspl_a_weighted"][observer]
        print(f"  Observer {observer}: {unweighted:.2f} dB, {a_weighted:.2f} dBA")


def print_oaspl_differences(
    dataset: AcousticDataset,
    simulated_dba_by_observer: dict[int, float],
    measured_dba_by_observer: dict[int, float],
    observers: tuple[int, ...],
) -> None:
    matched_observers = [
        observer
        for observer in observers
        if observer in simulated_dba_by_observer and observer in measured_dba_by_observer
    ]
    if not matched_observers:
        return

    print(f"A-weighted OASPL differences for {_case_label(dataset.source_file)} (simulation - measurement):")
    for observer in matched_observers:
        simulated = simulated_dba_by_observer[observer]
        measured = measured_dba_by_observer[observer]
        difference = simulated - measured
        if abs(difference) < 0.005:
            difference = 0.0
        print(
            f"  Observer {observer}: simulation {simulated:.2f} dBA, "
            f"measurement {measured:.2f} dBA, difference {difference:+.2f} dBA"
        )


def write_poc2_test_results(
    dataset: AcousticDataset,
    simulated_dba_by_observer: dict[int, float],
    measured_dba_by_observer: dict[int, float],
    output_dir: Path,
) -> Path:
    """Save measured POC2 OASPL and signed simulation errors."""
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / "poc2_test_results.csv"
    with output_file.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(("observer", "angle_deg", "simulation_dba", "experiment_dba", "simulation_minus_experiment_dba"))
        for observer in sorted(measured_dba_by_observer.keys() & simulated_dba_by_observer.keys()):
            simulated = simulated_dba_by_observer[observer]
            measured = measured_dba_by_observer[observer]
            difference = simulated - measured
            if abs(difference) < 0.005:
                difference = 0.0
            writer.writerow((observer, OBSERVER_ANGLES_DEG[observer], f"{simulated:.2f}", f"{measured:.2f}", f"{difference:+.2f}"))
    print(f"Wrote POC2 test results for {_legend_label(dataset.source_file)}: {output_file}")
    return output_file


def _case_label(source_file: Path) -> str:
    return re.split(r"_(?:results_)?total_acoustics", source_file.name, maxsplit=1)[0]


def _case_id(source_file: Path) -> str:
    match = CASE_ID_PATTERN.search(source_file.name)
    if match is None:
        raise ValueError(f"Could not infer case id from {source_file}")
    return match.group(0)


def _step_label(source_file: Path) -> str:
    for path_part in reversed(source_file.resolve().parts):
        path_part_lower = path_part.lower()
        if path_part_lower in CASE_TIMING_BY_PREFIX:
            return path_part_lower

    case_label = _case_label(source_file)
    for separator in ("_case-", "_"):
        if separator in case_label:
            return case_label.split(separator, 1)[0]
    return case_label


def _legend_label(source_file: Path) -> str:
    label = LEGEND_LABELS_BY_FILE.get(source_file.resolve())
    if label is not None:
        return label
    if LEGEND_LABEL_OVERRIDE is not None:
        return LEGEND_LABEL_OVERRIDE
    step_label = _step_label(source_file)
    path_parts = {part.lower() for part in source_file.resolve().parts}
    if "poc2x2" in path_parts:
        return "total - Case: Q40-LT"
    if "q40lt" in path_parts:
        return "total - Case: Q40LT"
    if "poc2" in path_parts:
        return f"total - Case: POC2 {step_label} (1000 RPM)"
    if "eve lifter4b" in path_parts or (
        CASE_ID_PATTERN.search(source_file.name)
        and _case_id(source_file) == "case-cf29e49e-da21-43ad-80e2-61a6f97bec80"
    ):
        return "total - Case: DUC EVE100 CCW Lifter4B"
    return f"total - Case: {step_label}"


def _surface_legend_label(surface_dataset: SurfaceResolvedDataset, case_index: int) -> str:
    group_label = f"zone_r1/{surface_dataset.group_name}"
    return f"{group_label} - {_legend_label(surface_dataset.case_source_file)}"


def _surface_resolved_dir_for_case(total_csv_file: Path) -> Path:
    return Path(total_csv_file).parent / "surface_resolved"


def _ensure_bunched_surface_files(surface_resolved_dir: Path) -> None:
    if any(surface_resolved_dir.glob("*_results_surface_zone_r1_blade1_acoustics_v*.csv")):
        return
    if bunch_surface_resolved_acoustics is None:
        return
    if any(surface_resolved_dir.glob("*_results_surface_zone_r1_blade1*_face*_acoustics_v*.csv")):
        bunch_surface_resolved_acoustics(surface_resolved_dir=surface_resolved_dir)


def _surface_file_for_case(total_csv_file: Path, group_name: str) -> Path | None:
    surface_resolved_dir = _surface_resolved_dir_for_case(total_csv_file)
    if not surface_resolved_dir.exists():
        return None
    _ensure_bunched_surface_files(surface_resolved_dir)
    case_id = _case_id(total_csv_file)
    matches = sorted(
        surface_resolved_dir.glob(f"{case_id}_results_surface_zone_r1_{group_name}_acoustics_v*.csv")
    )
    if not matches:
        return None
    if len(matches) > 1:
        raise ValueError(f"Found multiple surface files for {case_id} {group_name}: {matches}")
    return matches[0]


def _timing_for_file(csv_file: Path, rpm: float | None, time_steps_per_revolution: float | None) -> tuple[float, float]:
    if rpm is not None and time_steps_per_revolution is not None:
        return rpm, time_steps_per_revolution

    step_label = _step_label(csv_file)
    default_rpm, default_steps_per_revolution = CASE_TIMING_BY_PREFIX.get(
        step_label,
        (DEFAULT_RPM, DEFAULT_TIME_STEPS_PER_REVOLUTION),
    )
    if "poc2" in {part.lower() for part in csv_file.resolve().parts}:
        default_rpm = POC2_RPM
    elif "q40lt" in {part.lower() for part in csv_file.resolve().parts}:
        default_rpm = Q40LT_RPM
    return (
        default_rpm if rpm is None else rpm,
        default_steps_per_revolution if time_steps_per_revolution is None else time_steps_per_revolution,
    )


def _reference_values_for_file(
    csv_file: Path,
    reference_density_kg_m3: float | None,
    reference_velocity_m_s: float | None,
    reference_temperature_k: float | None,
    l_grid_unit_m: float | None,
) -> tuple[float, float, float, float]:
    if (
        reference_density_kg_m3 is not None
        and reference_velocity_m_s is not None
        and reference_temperature_k is not None
        and l_grid_unit_m is not None
    ):
        return reference_density_kg_m3, reference_velocity_m_s, reference_temperature_k, l_grid_unit_m

    path_parts = {part.lower() for part in csv_file.resolve().parts}
    if "q40lt" in path_parts:
        default_density = Q40LT_REFERENCE_DENSITY_KG_M3
        default_velocity = Q40LT_REFERENCE_VELOCITY_M_S
        default_temperature = Q40LT_REFERENCE_TEMPERATURE_K
        default_l_grid_unit_m = Q40LT_L_GRID_UNIT_M
    elif "poc2x2" in path_parts:
        default_density = POC2X2_REFERENCE_DENSITY_KG_M3
        default_velocity = POC2X2_REFERENCE_VELOCITY_M_S
        default_temperature = POC2X2_REFERENCE_TEMPERATURE_K
        default_l_grid_unit_m = POC2X2_L_GRID_UNIT_M
    elif "poc2" in path_parts:
        default_density = LOCAL_RESULTS_REFERENCE_DENSITY_KG_M3
        default_velocity = POC2_REFERENCE_VELOCITY_M_S
        default_temperature = LOCAL_RESULTS_REFERENCE_TEMPERATURE_K
        default_l_grid_unit_m = LOCAL_RESULTS_L_GRID_UNIT_M
    else:
        default_density = LOCAL_RESULTS_REFERENCE_DENSITY_KG_M3
        default_velocity = LOCAL_RESULTS_REFERENCE_VELOCITY_M_S
        default_temperature = LOCAL_RESULTS_REFERENCE_TEMPERATURE_K
        default_l_grid_unit_m = LOCAL_RESULTS_L_GRID_UNIT_M

    return (
        default_density if reference_density_kg_m3 is None else reference_density_kg_m3,
        default_velocity if reference_velocity_m_s is None else reference_velocity_m_s,
        default_temperature if reference_temperature_k is None else reference_temperature_k,
        default_l_grid_unit_m if l_grid_unit_m is None else l_grid_unit_m,
    )


def render_bar_plot(
    values_by_observer: dict[int, float],
    plot: PlotDefinition,
    source_file: Path,
    output_dir: Path = PLOTS_DIR,
    measurement_oaspl_dba_by_observer: dict[int, float] | None = None,
    measurement_label: str = MEASUREMENT_LABEL,
) -> Path:
    output_dir = _single_weighting_output_dir(output_dir, plot.name == "oaspl_a_weighted")
    observers = sorted(values_by_observer)
    values = [values_by_observer[observer] for observer in observers]

    fig, ax = plt.subplots(figsize=(8, 5))
    bars = ax.bar([str(observer) for observer in observers], values, color="#4477aa")
    ax.bar_label(bars, labels=[f"{value:.1f}" for value in values], padding=3, fontsize=9)
    all_values = list(values)
    if plot.name == "oaspl_a_weighted" and measurement_oaspl_dba_by_observer:
        measured_observers = [observer for observer in observers if observer in measurement_oaspl_dba_by_observer]
        if measured_observers:
            measured_values = [measurement_oaspl_dba_by_observer[observer] for observer in measured_observers]
            all_values.extend(measured_values)
            ax.plot(
                [observers.index(observer) for observer in measured_observers],
                measured_values,
                color="black",
                marker="x",
                markersize=8,
                linestyle="None",
                label=measurement_label,
            )
            ax.legend(loc="upper right")
    ax.set_xlabel("Observer")
    if "pusher" in {part.lower() for part in Path(source_file).parts}:
        ax.set_xticks(range(len(observers)), labels=[
            f"{observer}\n{angle}°" if (angle := _observer_angle_deg(source_file, observer)) is not None
            else str(observer)
            for observer in observers
        ])
    ax.set_ylabel(plot.ylabel)
    ax.set_title(f"{plot.title} - {_case_label(source_file)}")
    ax.grid(axis="y", alpha=0.3)
    y_min = min(all_values)
    y_max = max(all_values)
    padding = 3.0 if y_min == y_max else 0.15 * (y_max - y_min)
    ax.set_ylim(y_min - padding, y_max + padding)
    fig.tight_layout()

    output_file = output_dir / f"{_case_label(source_file)}_{plot.filename}"
    fig.savefig(output_file, dpi=200)
    plt.close(fig)
    return output_file


def render_comparison_plot(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    measurement_oaspl_dba_by_observer: dict[int, float] | None = None,
    measurement_label: str = MEASUREMENT_LABEL,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    fig, axes = plt.subplots(1, 2, figsize=(18, 5.5), sharex=True)
    markers = ("o", ">", "s", "D")

    for ax, plot in zip(axes, PLOTS):
        all_values = []
        for index, (dataset, values_by_plot) in enumerate(results_by_case):
            values_by_observer = values_by_plot[plot.name]
            values = [values_by_observer[observer] for observer in observers]
            all_values.extend(values)
            ax.plot(
                observers,
                values,
                linewidth=1.2,
                markersize=5,
                label=_legend_label(dataset.source_file),
                color=_line_color(dataset.source_file, index),
                marker=markers[index % len(markers)],
                linestyle="-",
            )

        if plot.name == "oaspl_a_weighted" and measurement_oaspl_dba_by_observer:
            measured_observers = [observer for observer in observers if observer in measurement_oaspl_dba_by_observer]
            if measured_observers:
                measured_values = [measurement_oaspl_dba_by_observer[observer] for observer in measured_observers]
                all_values.extend(measured_values)
                ax.plot(
                    measured_observers,
                    measured_values,
                    markersize=8,
                    label=measurement_label,
                    color="black",
                    marker="x",
                    linestyle="None",
                )

        y_min = min(all_values)
        y_max = max(all_values)
        padding = 0.2 * (y_max - y_min) if y_max > y_min else 3.0
        y_lower = y_min - max(padding, 8.0)
        y_upper = y_max + padding
        ax.set_ylim(y_lower, y_upper)

        angle_label_y = y_lower + 0.28 * (y_upper - y_lower)
        for observer in observers:
            angles = {_observer_angle_deg(dataset.source_file, observer) for dataset, _ in results_by_case}
            angle = next(iter(angles)) if len(angles) == 1 else None
            if angle is None:
                continue
            ax.vlines(observer, y_lower, angle_label_y - 0.02 * (y_upper - y_lower), colors="black", linewidth=1.2)
            ax.text(
                observer - 0.1,
                angle_label_y,
                f"{angle} deg",
                ha="left",
                va="bottom",
                fontsize=10,
            )

        ax.set_title(plot.title)
        ax.set_xlabel("Observer")
        ax.set_ylabel(plot.ylabel)
        ax.set_xlim(0, max(observers) + 1)
        ax.set_xticks(range(0, max(observers) + 2))
        ax.grid(True, alpha=0.3)
        ax.legend(loc="upper right")

    fig.tight_layout()
    output_file = output_dir / COMPARISON_OUTPUT_FILE
    fig.savefig(output_file, dpi=200)
    plt.close(fig)
    return output_file


def _frequency_tick_label(value: float, position: int) -> str:
    del position
    if value == 0:
        return "0"
    return f"{value:.0f}"


def _weighting_filename_part() -> str:
    return "_a_weighted" if APPLY_A_WEIGHTING_TO_SPECTRA else ""


def _single_weighting_output_dir(output_dir: Path, a_weighted: bool) -> Path:
    plot_dir = output_dir / ("a-weighted" if a_weighted else "unweighted")
    plot_dir.mkdir(parents=True, exist_ok=True)
    return plot_dir


def _style_frequency_axis(ax) -> None:
    ax.tick_params(axis="both", which="major", labelsize=13)
    ax.tick_params(axis="both", which="minor", labelsize=10)

    ax.grid(True, which="major", linestyle="--", alpha=0.6)
    ax.grid(True, which="minor", linestyle="--", alpha=0.35)

    ax.xaxis.set_major_locator(LogLocator(base=10))
    ax.xaxis.set_minor_locator(LogLocator(base=10.0, subs=(2, 4, 6, 8), numticks=5))
    ax.xaxis.set_major_formatter(FuncFormatter(_frequency_tick_label))
    ax.xaxis.set_minor_formatter(
        FuncFormatter(lambda value, position: "" if value == 8000 else _frequency_tick_label(value, position))
    )

    ax.tick_params(axis="x", which="major", labelsize=10)
    ax.tick_params(axis="x", which="minor", labelsize=8)
    ax.tick_params(axis="y", which="major", labelsize=10)
    ax.tick_params(axis="y", which="minor", labelsize=8)

    ax.yaxis.set_major_formatter(FuncFormatter(_frequency_tick_label))
    ax.yaxis.set_minor_locator(MultipleLocator(5))
    ax.yaxis.set_minor_formatter(FuncFormatter(_frequency_tick_label))


def load_poc2x_1000rpm_measurement_data(
    workbook_file: Path,
) -> Poc2xMeasurementData:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("openpyxl is required to load measurement Excel overlays.") from error

    workbook_file = Path(workbook_file)
    workbook = load_workbook(workbook_file, data_only=True, read_only=True)
    try:
        worksheet = workbook.worksheets[0]
        for observer, column_index in POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER.items():
            header_cell = worksheet.cell(row=3, column=column_index)
            expected = f"{OBSERVER_ANGLES_DEG[observer]} 1000 rpm"
            if expected not in str(header_cell.value or "").lower():
                raise ValueError(f"Expected {expected!r} in {worksheet.title}!{header_cell.coordinate}")

        oaspl_dba_by_observer = {}
        for observer, column_index in POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER.items():
            value = worksheet.cell(row=28, column=column_index).value
            if isinstance(value, (int, float)) and not isinstance(value, bool) and np.isfinite(value):
                oaspl_dba_by_observer[observer] = float(value)

        frequencies_by_observer = {
            observer: []
            for observer in POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER
        }
        spl_by_observer = {
            observer: []
            for observer in POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER
        }
        for row in worksheet.iter_rows(min_row=4, values_only=True):
            frequency_hz = row[0] if row else None
            if not isinstance(frequency_hz, (int, float)):
                continue
            for observer, column_index in POC2X_1000RPM_MEASUREMENT_COLUMNS_BY_OBSERVER.items():
                spl_db = row[column_index - 1]
                if isinstance(spl_db, (int, float)):
                    frequencies_by_observer[observer].append(float(frequency_hz))
                    spl_by_observer[observer].append(float(spl_db))

        measurement_spectra = {}
        for observer, frequencies_hz in frequencies_by_observer.items():
            if not frequencies_hz:
                continue
            measurement_spectra[observer] = MeasurementSpectrum(
                frequencies_hz=np.asarray(frequencies_hz, dtype=float),
                spl_db=np.asarray(spl_by_observer[observer], dtype=float),
            )
        return Poc2xMeasurementData(
            spectra_by_observer=measurement_spectra,
            oaspl_dba_by_observer=oaspl_dba_by_observer,
        )
    finally:
        workbook.close()


def load_poc2x_1000rpm_measurement_spectra(
    workbook_file: Path,
) -> dict[int, MeasurementSpectrum]:
    return load_poc2x_1000rpm_measurement_data(workbook_file).spectra_by_observer


def render_spectrum_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
) -> list[Path]:
    output_dir = _single_weighting_output_dir(output_dir, APPLY_A_WEIGHTING_TO_SPECTRA)
    output_files = []

    for observer in observers:
        fig, ax = plt.subplots(figsize=(8.5, 5.2))
        for index, (dataset, _) in enumerate(results_by_case):
            pressure_pa = dataset.pressure_by_observer[observer] * dataset.pressure_scale_pa
            frequencies, spl_db = spl_spectrum_db(pressure_pa, dataset.sample_rate_hz)
            frequency_mask = (
                (frequencies >= SPECTRUM_MIN_FREQUENCY_HZ)
                & (frequencies <= SPECTRUM_MAX_FREQUENCY_HZ)
            )
            ax.plot(
                frequencies[frequency_mask],
                spl_db[frequency_mask],
                label=_legend_label(dataset.source_file),
                color=_line_color(dataset.source_file, index),
                linestyle="-",
                linewidth=1.2,
            )

        ax.set_title(f"Spectra for observer ID: {observer}")
        ax.set_xlabel("Frequency [Hz]")
        ax.set_ylabel("SPL [dBA]" if APPLY_A_WEIGHTING_TO_SPECTRA else "SPL [dB]")
        ax.set_xscale("log")
        ax.set_xlim(SPECTRUM_MIN_FREQUENCY_HZ, SPECTRUM_MAX_FREQUENCY_HZ)
        ax.set_ylim(*SPECTRUM_Y_LIMITS_DB)
        _style_frequency_axis(ax)
        ax.legend(loc="upper right", fontsize=12)
        fig.tight_layout()

        output_file = output_dir / f"spectrum{_weighting_filename_part()}_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)
    return output_files


def render_third_octave_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    measurement_spectra_by_observer: dict[int, MeasurementSpectrum] | None = None,
    measurement_label: str = MEASUREMENT_LABEL,
) -> list[Path]:
    output_dir = _single_weighting_output_dir(output_dir, APPLY_A_WEIGHTING_TO_SPECTRA)
    output_files = []
    markers = ("o", "s", "D", "^")

    for observer in observers:
        fig, ax = plt.subplots(figsize=(8.5, 5.2))
        upper_frequency_limit = SPECTRUM_MAX_FREQUENCY_HZ
        for index, (dataset, _) in enumerate(results_by_case):
            pressure_pa = dataset.pressure_by_observer[observer] * dataset.pressure_scale_pa
            center_frequencies, band_spl_db = third_octave_spectrum_db(pressure_pa, dataset.sample_rate_hz)
            upper_frequency_limit = max(upper_frequency_limit, float(np.max(center_frequencies)))
            valid_band_mask = np.isfinite(band_spl_db)
            ax.plot(
                center_frequencies[valid_band_mask],
                band_spl_db[valid_band_mask],
                label=_legend_label(dataset.source_file),
                color=_line_color(dataset.source_file, index),
                marker=markers[index % len(markers)],
                markerfacecolor="none",
                linestyle="-",
                linewidth=1.5,
            )

        if measurement_spectra_by_observer is not None:
            measurement_spectrum = measurement_spectra_by_observer.get(observer)
            if measurement_spectrum is not None:
                frequency_mask = (
                    (measurement_spectrum.frequencies_hz >= SPECTRUM_MIN_FREQUENCY_HZ)
                    & (measurement_spectrum.frequencies_hz <= SPECTRUM_MAX_FREQUENCY_HZ)
                )
                if np.any(frequency_mask):
                    upper_frequency_limit = max(
                        upper_frequency_limit,
                        float(np.max(measurement_spectrum.frequencies_hz[frequency_mask])),
                    )
                    ax.plot(
                        measurement_spectrum.frequencies_hz[frequency_mask],
                        measurement_spectrum.spl_db[frequency_mask],
                        label=measurement_label,
                        color="black",
                        linestyle="--",
                        marker="x",
                        markersize=6,
                        linewidth=2.0,
                    )

        ax.set_title(f"1/3 Octave Spectra for observer ID: {observer}")
        ax.set_xlabel("Frequency [Hz]")
        ax.set_ylabel(r"SPL$_{1/3}$ [dBA]" if APPLY_A_WEIGHTING_TO_SPECTRA else r"SPL$_{1/3}$ [dB]")
        ax.set_xscale("log")
        ax.set_xlim(SPECTRUM_MIN_FREQUENCY_HZ, upper_frequency_limit)
        ax.set_ylim(*SPECTRUM_Y_LIMITS_DB)
        _style_frequency_axis(ax)
        ax.legend(loc="lower right", fontsize=10)
        fig.tight_layout()

        output_file = output_dir / f"third_octave_spectrum{_weighting_filename_part()}_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)
    return output_files


def _validate_matching_surface_signals(
    reference_dataset: AcousticDataset,
    other_dataset: AcousticDataset,
    pair_label: str,
) -> None:
    if len(reference_dataset.time) != len(other_dataset.time):
        raise ValueError(
            f"Cannot sum {pair_label}: {reference_dataset.source_file.name} and "
            f"{other_dataset.source_file.name} have different sample counts."
        )
    if not np.allclose(reference_dataset.time, other_dataset.time, rtol=0.0, atol=1.0e-12):
        raise ValueError(
            f"Cannot sum {pair_label}: {reference_dataset.source_file.name} and "
            f"{other_dataset.source_file.name} have different time coordinates."
        )


def render_surface_pair_third_octave_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    group_pairs: tuple[tuple[str, tuple[str, ...]], ...] = SURFACE_THIRD_OCTAVE_GROUP_PAIRS,
) -> list[Path]:
    output_dir = _single_weighting_output_dir(output_dir, APPLY_A_WEIGHTING_TO_SPECTRA)
    output_files = []
    pair_groups = tuple(group for _, groups in group_pairs for group in groups)
    surface_datasets = load_surface_resolved_datasets(
        results_by_case,
        groups=pair_groups,
        observers=observers,
    )
    if not surface_datasets:
        return []

    surface_datasets_by_case: dict[Path, dict[str, SurfaceResolvedDataset]] = {}
    for surface_dataset in surface_datasets:
        surface_datasets_by_case.setdefault(surface_dataset.case_source_file, {})[
            surface_dataset.group_name
        ] = surface_dataset

    pair_styles = {
        "blade1+blade3": {"linestyle": "-", "marker": "o"},
        "blade2+blade4": {"linestyle": "--", "marker": "s"},
    }

    for observer in observers:
        fig, ax = plt.subplots(figsize=(8.5, 5.2))
        upper_frequency_limit = SPECTRUM_MAX_FREQUENCY_HZ
        for case_index, (total_dataset, _) in enumerate(results_by_case):
            case_surfaces = surface_datasets_by_case.get(total_dataset.source_file, {})
            for pair_label, group_names in group_pairs:
                missing_groups = [group_name for group_name in group_names if group_name not in case_surfaces]
                if missing_groups:
                    print(
                        f"Skipping {pair_label} for {total_dataset.source_file.name}; "
                        f"missing {missing_groups}."
                    )
                    continue

                first_surface_dataset = case_surfaces[group_names[0]].dataset
                summed_pressure = np.zeros_like(first_surface_dataset.pressure_by_observer[observer])
                for group_name in group_names:
                    surface_dataset = case_surfaces[group_name].dataset
                    _validate_matching_surface_signals(first_surface_dataset, surface_dataset, pair_label)
                    summed_pressure += surface_dataset.pressure_by_observer[observer]

                pressure_pa = summed_pressure * total_dataset.pressure_scale_pa
                center_frequencies, band_spl_db = third_octave_spectrum_db(
                    pressure_pa,
                    total_dataset.sample_rate_hz,
                )
                upper_frequency_limit = max(upper_frequency_limit, float(np.max(center_frequencies)))
                valid_band_mask = np.isfinite(band_spl_db)
                style = pair_styles.get(
                    pair_label,
                    {
                        "linestyle": "-",
                        "marker": "o",
                    },
                )
                ax.plot(
                    center_frequencies[valid_band_mask],
                    band_spl_db[valid_band_mask],
                    label=f"{pair_label} - {_legend_label(total_dataset.source_file)}",
                    color=_line_color(total_dataset.source_file, case_index),
                    markerfacecolor="none",
                    linewidth=1.2,
                    **style,
                )

        ax.set_title(f"Blade Pair 1/3 Octave Spectra for observer ID: {observer}")
        ax.set_xlabel("Frequency [Hz]")
        ax.set_ylabel(r"SPL$_{1/3}$ [dBA]" if APPLY_A_WEIGHTING_TO_SPECTRA else r"SPL$_{1/3}$ [dB]")
        ax.set_xscale("log")
        ax.set_xlim(SPECTRUM_MIN_FREQUENCY_HZ, upper_frequency_limit)
        ax.set_ylim(*SPECTRUM_Y_LIMITS_DB)
        _style_frequency_axis(ax)
        ax.legend(loc="upper right", fontsize=12)
        fig.tight_layout()

        output_file = output_dir / f"surface_pair_third_octave_spectrum{_weighting_filename_part()}_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)

    return output_files


def _summed_surface_pair_pressure(
    case_surfaces: dict[str, SurfaceResolvedDataset],
    group_names: tuple[str, ...],
    observer: int,
    pair_label: str,
) -> np.ndarray:
    first_surface_dataset = case_surfaces[group_names[0]].dataset
    summed_pressure = np.zeros_like(first_surface_dataset.pressure_by_observer[observer])
    for group_name in group_names:
        surface_dataset = case_surfaces[group_name].dataset
        _validate_matching_surface_signals(first_surface_dataset, surface_dataset, pair_label)
        summed_pressure += surface_dataset.pressure_by_observer[observer]
    return summed_pressure


def render_surface_pair_oaspl_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    group_pairs: tuple[tuple[str, tuple[str, ...]], ...] = SURFACE_THIRD_OCTAVE_GROUP_PAIRS,
) -> Path | None:
    output_dir.mkdir(parents=True, exist_ok=True)
    pair_groups = tuple(group for _, groups in group_pairs for group in groups)
    surface_datasets = load_surface_resolved_datasets(
        results_by_case,
        groups=pair_groups,
        observers=observers,
    )
    if not surface_datasets:
        return None

    surface_datasets_by_case: dict[Path, dict[str, SurfaceResolvedDataset]] = {}
    for surface_dataset in surface_datasets:
        surface_datasets_by_case.setdefault(surface_dataset.case_source_file, {})[
            surface_dataset.group_name
        ] = surface_dataset

    fig, axes = plt.subplots(1, 2, figsize=(18, 5.5), sharex=True)
    plot_definitions = (
        ("Unweighted Surface-Pair OASPL", "OASPL [dB]", oaspl_unweighted_db),
        ("A-weighted Surface-Pair OASPL", "A-weighted OASPL [dBA]", oaspl_a_weighted_db),
    )
    pair_styles = {
        "blade1+blade3": {"linestyle": "-", "marker": "o"},
        "blade2+blade4": {"linestyle": "--", "marker": "s"},
    }

    for ax, (title, ylabel, evaluator) in zip(axes, plot_definitions):
        all_values = []
        for case_index, (total_dataset, _) in enumerate(results_by_case):
            case_surfaces = surface_datasets_by_case.get(total_dataset.source_file, {})
            for pair_label, group_names in group_pairs:
                missing_groups = [group_name for group_name in group_names if group_name not in case_surfaces]
                if missing_groups:
                    print(
                        f"Skipping {pair_label} OASPL for {total_dataset.source_file.name}; "
                        f"missing {missing_groups}."
                    )
                    continue

                values = []
                for observer in observers:
                    summed_pressure = _summed_surface_pair_pressure(
                        case_surfaces,
                        group_names,
                        observer,
                        pair_label,
                    )
                    pressure_pa = summed_pressure * total_dataset.pressure_scale_pa
                    values.append(evaluator(pressure_pa, total_dataset.sample_rate_hz))
                all_values.extend(values)
                style = pair_styles.get(pair_label, {"linestyle": "-", "marker": "o"})
                ax.plot(
                    observers,
                    values,
                    label=f"{pair_label} - {_legend_label(total_dataset.source_file)}",
                    color=_line_color(total_dataset.source_file, case_index),
                    linewidth=1.2,
                    markersize=5,
                    markerfacecolor="none",
                    **style,
                )

        if all_values:
            y_min = min(all_values)
            y_max = max(all_values)
            padding = max(3.0, 0.2 * (y_max - y_min)) if y_max > y_min else 3.0
            ax.set_ylim(y_min - padding, y_max + padding)
        ax.set_title(title)
        ax.set_xlabel("Observer")
        ax.set_ylabel(ylabel)
        ax.set_xlim(0, max(observers) + 1)
        ax.set_xticks(range(0, max(observers) + 2))
        ax.grid(True, alpha=0.3)
        ax.legend(loc="upper right")

    fig.tight_layout()
    output_file = output_dir / "surface_pair_oaspl_observers_1_5.png"
    fig.savefig(output_file, dpi=200)
    plt.close(fig)
    return output_file


def _bpf_frequency_hz(dataset: AcousticDataset, harmonic: int, blade_count: int = DEFAULT_BLADE_COUNT) -> float:
    return harmonic * blade_count * dataset.rpm / 60.0


def render_bpf_component_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    blade_count: int = DEFAULT_BLADE_COUNT,
) -> list[Path]:
    output_dir = _single_weighting_output_dir(output_dir, False)
    output_files = []
    component_definitions = (
        ("Thickness pressure", "thickness_by_observer"),
        ("Loading pressure", "loading_by_observer"),
        ("Total pressure", "pressure_by_observer"),
    )
    markers = ("^", ">", "s", "o")

    for harmonic, harmonic_label in ((1, "first"), (2, "second")):
        fig, axes = plt.subplots(1, 3, figsize=(16, 5.8), sharey=True)
        all_component_values = []
        actual_frequencies = []

        for ax, (component_title, attribute_name) in zip(axes, component_definitions):
            for index, (dataset, _) in enumerate(results_by_case):
                target_frequency = _bpf_frequency_hz(dataset, harmonic, blade_count)
                component_by_observer = getattr(dataset, attribute_name)
                values = []
                nearest_frequencies = []
                for observer in observers:
                    component_pressure = component_by_observer[observer] * dataset.pressure_scale_pa
                    nearest_frequency, spl_db = tone_spl_db(
                        component_pressure,
                        dataset.sample_rate_hz,
                        target_frequency,
                    )
                    values.append(spl_db)
                    nearest_frequencies.append(nearest_frequency)
                all_component_values.extend(values)
                actual_frequencies.append(float(np.mean(nearest_frequencies)))
                ax.plot(
                    values,
                    observers,
                    label=_legend_label(dataset.source_file),
                    color=_line_color(dataset.source_file, index),
                    marker=markers[index % len(markers)],
                    markerfacecolor=_line_color(dataset.source_file, index),
                    linestyle="-",
                    linewidth=1.2,
                )

            ax.set_title(component_title)
            ax.set_xlabel("SPL [dB]")
            ax.set_ylim(0, max(observers) + 1)
            ax.set_yticks(range(0, max(observers) + 2))
            ax.grid(True, alpha=0.3)

        axes[0].set_ylabel("Observer ID")
        axes[-1].legend(loc="lower left")
        if all_component_values:
            x_min = min(all_component_values)
            x_max = max(all_component_values)
            x_padding = max(4.0, 0.15 * (x_max - x_min))
            for ax in axes:
                ax.set_xlim(x_min - x_padding, x_max + x_padding)

        average_frequency = float(np.mean(actual_frequencies)) if actual_frequencies else 0.0
        fig.suptitle(
            f"Thickness and Loading Component for the {harmonic_label.title()} BPF "
            f"({average_frequency:.1f} Hz)",
            y=0.98,
        )
        fig.tight_layout()

        output_file = output_dir / f"thickness_loading_{harmonic_label}_bpf.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)

    return output_files


def render_harmonic_content_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    blade_count: int = DEFAULT_BLADE_COUNT,
) -> list[Path]:
    output_dir = _single_weighting_output_dir(output_dir, False)
    output_files = []
    markers = ("^", ">", "s", "o")

    harmonics = np.asarray(HARMONIC_INDICES, dtype=int)
    for observer in observers:
        fig, ax = plt.subplots(figsize=(10.5, 6.0))
        for index, (dataset, _) in enumerate(results_by_case):
            bpf_frequency = _bpf_frequency_hz(dataset, 1, blade_count)
            harmonic_levels = []
            for harmonic in harmonics:
                target_frequency = float(harmonic * bpf_frequency)
                pressure_pa = dataset.pressure_by_observer[observer] * dataset.pressure_scale_pa
                spl_db = harmonic_third_octave_spl_db(
                    pressure_pa,
                    dataset.sample_rate_hz,
                    target_frequency,
                )
                harmonic_levels.append(spl_db)

            color = _line_color(dataset.source_file, index)
            ax.vlines(
                harmonics,
                0.0,
                harmonic_levels,
                colors=color,
                linewidth=4.0,
                alpha=0.95,
            )
            ax.plot(
                harmonics,
                harmonic_levels,
                linestyle="none",
                markersize=8,
                label=_legend_label(dataset.source_file),
                color=color,
                marker=markers[index % len(markers)],
                markerfacecolor=color,
            )

        ax.set_title(f"Harmonic content for observer ID: {observer}")
        ax.set_xlabel("Harmonics")
        ax.set_ylabel(r"SPL$_{1/3}$ [dB]")
        ax.set_xlim(0, max(HARMONIC_INDICES) + 1)
        ax.set_xticks(range(0, max(HARMONIC_INDICES) + 2))
        ax.set_ylim(*HARMONIC_Y_LIMITS_DB)
        ax.grid(True, alpha=0.3)
        ax.legend(loc="upper right")
        fig.tight_layout()

        output_file = output_dir / f"harmonic_content_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)

    return output_files


def render_pressure_time_history_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_files = []
    for observer in observers:
        fig, ax = plt.subplots(figsize=(10.5, 6.0))
        all_pressure_values = []

        for index, (dataset, _) in enumerate(results_by_case):
            time_seconds, pressure_pa = _pressure_time_history(dataset, observer)
            all_pressure_values.extend(pressure_pa)
            ax.plot(
                time_seconds,
                pressure_pa,
                label=_legend_label(dataset.source_file),
                color=_line_color(dataset.source_file, index),
                linestyle="-",
                linewidth=1.2,
            )

        ax.set_title(f"Pressure-time History for observer ID: {observer}")
        ax.set_xlabel("Elapsed time [s]")
        ax.set_ylabel("Pressure [Pa]")
        ax.margins(x=0)
        if all_pressure_values:
            pressure_min = min(all_pressure_values)
            pressure_max = max(all_pressure_values)
            padding = 0.15 * (pressure_max - pressure_min) if pressure_max > pressure_min else 0.01
            ax.set_ylim(pressure_min - padding, pressure_max + padding)
        ax.grid(True, alpha=0.3)
        ax.legend(loc="upper left")
        fig.tight_layout()

        output_file = output_dir / f"pressure_time_history_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)

    return output_files


def render_surface_pressure_time_history_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    groups: tuple[str, ...] = SURFACE_TIME_HISTORY_GROUPS,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    surface_datasets = load_surface_resolved_datasets(
        results_by_case,
        groups=groups,
        observers=observers,
    )
    if not surface_datasets:
        return []

    output_files = []
    line_styles = {
        ("blade1", 0): {"color": "green", "linewidth": 1.2},
        ("blade2", 0): {"color": "red", "linewidth": 1.2},
        ("blade1", 1): {"color": "blue", "linewidth": 1.2},
        ("blade2", 1): {"color": "cyan", "linewidth": 1.2},
    }

    for observer in observers:
        fig, ax = plt.subplots(figsize=(10.5, 6.0))
        all_pressure_values = []
        case_indices: dict[Path, int] = {}

        for surface_dataset in surface_datasets:
            case_index = case_indices.setdefault(
                surface_dataset.case_source_file,
                len(case_indices),
            )
            dataset = surface_dataset.dataset
            time_seconds, pressure_pa = _pressure_time_history(dataset, observer)
            all_pressure_values.extend(pressure_pa)
            style = line_styles.get(
                (surface_dataset.group_name, case_index),
                {"color": _line_color(surface_dataset.case_source_file, case_index), "linewidth": 1.2},
            )
            ax.plot(
                time_seconds,
                pressure_pa,
                label=_surface_legend_label(surface_dataset, case_index),
                linestyle="-",
                **style,
            )

        ax.set_title(f"Pressure-time History for observer ID: {observer}")
        ax.set_xlabel("Elapsed time [s]")
        ax.set_ylabel("Pressure [Pa]")
        ax.margins(x=0)
        if all_pressure_values:
            pressure_min = min(all_pressure_values)
            pressure_max = max(all_pressure_values)
            padding = 0.15 * (pressure_max - pressure_min) if pressure_max > pressure_min else 0.01
            ax.set_ylim(pressure_min - padding, pressure_max + padding)
        ax.grid(True, alpha=0.3)
        ax.legend(loc="lower right")
        fig.tight_layout()

        output_file = output_dir / f"surface_pressure_time_history_blade1_blade2_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)

    return output_files


def _average_revolutions(signal: np.ndarray, revolution_count: int = 3) -> np.ndarray:
    signal = np.asarray(signal, dtype=float)
    samples_per_revolution = len(signal) // revolution_count
    if samples_per_revolution < 2:
        raise ValueError(
            f"Need at least {2 * revolution_count} samples to average {revolution_count} revolutions."
        )
    trimmed_sample_count = samples_per_revolution * revolution_count
    revolutions = signal[:trimmed_sample_count].reshape(revolution_count, samples_per_revolution)
    return np.mean(revolutions, axis=0)


def _average_revolution_count(dataset: AcousticDataset, default_revolution_count: int) -> int:
    if dataset.revolution_range is not None:
        return dataset.revolution_range.count
    return default_revolution_count


def render_average_pressure_time_history_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
    revolution_count: int = 3,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_files = []

    for observer in observers:
        fig, ax = plt.subplots(figsize=(10.5, 6.0))
        all_pressure_values = []

        for index, (dataset, _) in enumerate(results_by_case):
            pressure_pa = dataset.pressure_by_observer[observer] * dataset.pressure_scale_pa
            pressure_pa = _filter_zero_signal_edges(pressure_pa)
            samples_per_revolution = int(round(dataset.sample_rate_hz * 60.0 / dataset.rpm))
            complete_revolutions = len(pressure_pa) // samples_per_revolution
            average_count = min(
                _average_revolution_count(dataset, revolution_count), complete_revolutions,
            )
            if average_count < 1:
                raise ValueError(f"Observer {observer} has less than one complete acoustic revolution.")
            averaged_pressure_pa = _average_revolutions(
                pressure_pa[:average_count * samples_per_revolution],
                revolution_count=average_count,
            )
            all_pressure_values.extend(averaged_pressure_pa)
            time_seconds = np.arange(len(averaged_pressure_pa)) * dataset.sample_spacing_s

            ax.plot(
                time_seconds,
                averaged_pressure_pa,
                label=_legend_label(dataset.source_file),
                color=_line_color(dataset.source_file, index),
                linestyle="-",
                linewidth=1.2,
            )

        ax.set_title(f"Average Pressure-time History for observer ID: {observer}")
        ax.set_xlabel("Elapsed time [s]")
        ax.set_ylabel("Pressure [Pa]")
        ax.margins(x=0)
        if all_pressure_values:
            pressure_min = min(all_pressure_values)
            pressure_max = max(all_pressure_values)
            padding = 0.15 * (pressure_max - pressure_min) if pressure_max > pressure_min else 0.01
            ax.set_ylim(pressure_min - padding, pressure_max + padding)
        ax.grid(True, alpha=0.3)
        ax.legend(loc="upper left")
        fig.tight_layout()

        output_file = output_dir / f"average_pressure_time_history_observer_{observer}.png"
        fig.savefig(output_file, dpi=200)
        plt.close(fig)
        output_files.append(output_file)

    return output_files


def _spectrogram_spl_db(
    pressure_pa: np.ndarray,
    sample_rate_hz: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    pressure_pa = _filter_zero_signal_edges(pressure_pa)
    nperseg = min(SPECTROGRAM_NPERSEG, len(pressure_pa))
    if nperseg < 8:
        raise ValueError("Too few samples for spectrogram calculation.")

    frequencies, times, power_by_time = spectrogram(
        _fluctuating_pressure(pressure_pa),
        fs=sample_rate_hz,
        window=WINDOW,
        nperseg=nperseg,
        noverlap=nperseg // 2,
        scaling="spectrum",
        mode="psd",
    )
    spl_db = _db_from_power(power_by_time)
    return frequencies, times, spl_db


def render_spectrogram_plots(
    results_by_case: list[tuple[AcousticDataset, dict[str, dict[int, float]]]],
    observers: tuple[int, ...],
    output_dir: Path = PLOTS_DIR,
) -> list[Path]:
    output_dir = _single_weighting_output_dir(output_dir, False)
    output_files = []
    component_definitions = (
        ("PRESSURE", "pressure_by_observer", "total_pressure"),
        ("LOADING", "loading_by_observer", "loading_pressure"),
    )

    for case_index, (dataset, _) in enumerate(results_by_case):
        if len(results_by_case) == 1:
            case_label = _case_label(dataset.source_file)
        else:
            label_slug = re.sub(r"[^a-z0-9]+", "_", _legend_label(dataset.source_file).lower()).strip("_")
            case_label = f"{case_index + 1}_{label_slug}"
        for observer in observers:
            for component_title, attribute_name, filename_part in component_definitions:
                component_by_observer = getattr(dataset, attribute_name)
                pressure_pa = component_by_observer[observer] * dataset.pressure_scale_pa
                frequencies, times, spl_db = _spectrogram_spl_db(
                    pressure_pa,
                    dataset.sample_rate_hz,
                )

                frequency_mask = frequencies <= SPECTROGRAM_MAX_FREQUENCY_HZ
                frequencies = frequencies[frequency_mask]
                spl_db = spl_db[frequency_mask, :]
                time_fraction = times / times[-1] if times[-1] > 0.0 else times

                fig, ax = plt.subplots(figsize=(10.5, 6.0))
                mesh = ax.pcolormesh(
                    time_fraction,
                    frequencies,
                    spl_db,
                    shading="nearest",
                    cmap="jet",
                    vmin=SPECTROGRAM_LIMITS_DB[0],
                    vmax=SPECTROGRAM_LIMITS_DB[1],
                )
                colorbar = fig.colorbar(mesh, ax=ax, pad=0.10)
                colorbar.set_label("SPL [dB]")

                ax.set_title(f"{component_title} - Spectrogram for Observer ID: {observer}")
                ax.set_xlabel("Time Fraction")
                ax.set_ylabel("Frequency [Hz]")
                ax.set_xlim(0.0, 1.0)
                ax.set_ylim(0.0, SPECTROGRAM_MAX_FREQUENCY_HZ)
                ax.grid(True, color="white", linewidth=0.5, alpha=0.5)
                fig.tight_layout()

                output_file = output_dir / (
                    f"spectrogram_{filename_part}_{case_label}_observer_{observer}.png"
                )
                fig.savefig(output_file, dpi=200)
                plt.close(fig)
                output_files.append(output_file)

    return output_files


def make_plots(
    csv_file: Path | None = None,
    observers: tuple[int, ...] = DEFAULT_OBSERVERS,
    plots: tuple[PlotDefinition, ...] = PLOTS,
    output_dir: Path = PLOTS_DIR,
    rpm: float | None = None,
    time_steps_per_revolution: float | None = None,
    reference_density_kg_m3: float | None = None,
    reference_velocity_m_s: float | None = None,
    reference_temperature_k: float | None = None,
    l_grid_unit_m: float | None = None,
    revolution_range: RevolutionRange | None = None,
    measurement_data: Poc2xMeasurementData | None = None,
    measurement_label: str = MEASUREMENT_LABEL,
    csv_files: list[Path] | None = None,
    n_timesteps_skip: int = 0,
    skip_unit: str = "cfd",
) -> list[Path]:
    resolved_csv_file = (
        Path(csv_files[0]) if csv_files
        else _default_acoustic_files()[-1] if csv_file is None else Path(csv_file)
    )
    rpm, time_steps_per_revolution = _timing_for_file(
        resolved_csv_file,
        rpm,
        time_steps_per_revolution,
    )
    reference_density_kg_m3, reference_velocity_m_s, reference_temperature_k, l_grid_unit_m = _reference_values_for_file(
        resolved_csv_file,
        reference_density_kg_m3,
        reference_velocity_m_s,
        reference_temperature_k,
        l_grid_unit_m,
    )
    sources = csv_files or [resolved_csv_file]
    datasets = []
    for source in sources:
        source_rpm, source_steps = _timing_for_file(source, rpm, time_steps_per_revolution)
        references = _reference_values_for_file(
            source, reference_density_kg_m3, reference_velocity_m_s,
            reference_temperature_k, l_grid_unit_m,
        )
        datasets.append(load_acoustic_dataset(
            source, observers, rpm=source_rpm, time_steps_per_revolution=source_steps,
            reference_density_kg_m3=references[0], reference_velocity_m_s=references[1],
            reference_temperature_k=references[2], l_grid_unit_m=references[3],
            revolution_range=revolution_range if not csv_files and not n_timesteps_skip else None,
        ))
    dataset = combine_acoustic_datasets(datasets) if csv_files else datasets[0]
    if csv_files or n_timesteps_skip:
        print("Acoustic sources: " + ", ".join(str(source) for source in sources))
        dataset = prepare_acoustic_dataset(dataset, n_timesteps_skip, skip_unit, revolution_range)
    print(
        f"Loaded {dataset.source_file.name}: "
        f"{len(dataset.time)} samples, "
        f"revolutions={dataset.revolution_range.label if dataset.revolution_range is not None else 'all'}, "
        f"rpm={dataset.rpm:.3f}, "
        f"steps/rev={dataset.time_steps_per_revolution:g}, "
        f"sample_rate={dataset.sample_rate_hz:.3f} Hz, "
        f"rho={dataset.reference_density_kg_m3:.6f} kg/m^3, "
        f"T={dataset.reference_temperature_k:.3f} K, "
        f"a={dataset.speed_of_sound_m_s:.6f} m/s, "
        f"L_grid_unit={dataset.l_grid_unit_m:g} m, "
        f"U_ref={dataset.reference_velocity_m_s:.6f} m/s, "
        f"pressure_scale=gamma*p_inf={dataset.pressure_scale_pa:.3f} Pa"
    )

    output_files = []
    values_by_plot = {}
    for plot in plots:
        values_by_observer = evaluate_plot(dataset, plot)
        values_by_plot[plot.name] = values_by_observer
        output_file = render_bar_plot(
            values_by_observer,
            plot,
            dataset.source_file,
            output_dir,
            measurement_oaspl_dba_by_observer=(
                measurement_data.oaspl_dba_by_observer if measurement_data is not None else None
            ),
            measurement_label=measurement_label,
        )
        output_files.append(output_file)
        print(f"Wrote {output_file}")
    if all(plot.name in values_by_plot for plot in PLOTS):
        print_oaspl_values(dataset, values_by_plot, observers)
    if csv_files or n_timesteps_skip:
        output_dir.mkdir(parents=True, exist_ok=True)
        skipped_samples = (
            n_timesteps_skip if skip_unit == "acoustic"
            else int(round(n_timesteps_skip * 60.0 / (
                dataset.rpm * dataset.time_steps_per_revolution * dataset.sample_spacing_s
            )))
        )
        observer_windows = {}
        for observer in observers:
            elapsed, _ = _pressure_time_history(dataset, observer)
            observer_windows[str(observer)] = {
                "retained_samples": len(elapsed),
                "elapsed_span_s": float(elapsed[-1]),
                "skipped_acoustic_samples": skipped_samples,
            }
        manifest = {
            "sources": [str(source.resolve()) for source in sources],
            "n_timesteps_skip": n_timesteps_skip, "skip_unit": skip_unit,
            "sample_spacing_s": dataset.sample_spacing_s,
            "rpm": dataset.rpm, "solver_steps_per_revolution": dataset.time_steps_per_revolution,
            "density_kg_m3": dataset.reference_density_kg_m3,
            "temperature_k": dataset.reference_temperature_k,
            "reference_velocity_m_s": dataset.reference_velocity_m_s,
            "grid_unit_m": dataset.l_grid_unit_m,
            "observers": observer_windows,
            "oaspl": values_by_plot,
        }
        (output_dir / "analysis_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    if measurement_data is not None and "oaspl_a_weighted" in values_by_plot:
        print_oaspl_differences(
            dataset,
            values_by_plot["oaspl_a_weighted"],
            measurement_data.oaspl_dba_by_observer,
            observers,
        )
        if "poc2" in {part.lower() for part in dataset.source_file.resolve().parts}:
            write_poc2_test_results(
                dataset,
                values_by_plot["oaspl_a_weighted"],
                measurement_data.oaspl_dba_by_observer,
                output_dir,
            )
    for spectrum_output_file in render_spectrum_plots([(dataset, values_by_plot)], observers, output_dir):
        output_files.append(spectrum_output_file)
        print(f"Wrote {spectrum_output_file}")
    for harmonic_output_file in render_harmonic_content_plots(
        [(dataset, values_by_plot)], observers, output_dir,
    ):
        output_files.append(harmonic_output_file)
        print(f"Wrote {harmonic_output_file}")
    for third_octave_output_file in render_third_octave_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
        measurement_spectra_by_observer=(
            measurement_data.spectra_by_observer
            if measurement_data is not None and APPLY_A_WEIGHTING_TO_SPECTRA
            else None
        ),
        measurement_label=measurement_label,
    ):
        output_files.append(third_octave_output_file)
        print(f"Wrote {third_octave_output_file}")
    for pressure_time_output_file in render_pressure_time_history_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
    ):
        output_files.append(pressure_time_output_file)
        print(f"Wrote {pressure_time_output_file}")
    for surface_pressure_time_output_file in render_surface_pressure_time_history_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
    ):
        output_files.append(surface_pressure_time_output_file)
        print(f"Wrote {surface_pressure_time_output_file}")
    for surface_pair_third_octave_output_file in render_surface_pair_third_octave_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
    ):
        output_files.append(surface_pair_third_octave_output_file)
        print(f"Wrote {surface_pair_third_octave_output_file}")
    surface_pair_oaspl_output_file = render_surface_pair_oaspl_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
    )
    if surface_pair_oaspl_output_file is not None:
        output_files.append(surface_pair_oaspl_output_file)
        print(f"Wrote {surface_pair_oaspl_output_file}")
    for average_pressure_time_output_file in render_average_pressure_time_history_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
    ):
        output_files.append(average_pressure_time_output_file)
        print(f"Wrote {average_pressure_time_output_file}")
    for spectrogram_output_file in render_spectrogram_plots(
        [(dataset, values_by_plot)],
        observers,
        output_dir,
    ):
        output_files.append(spectrogram_output_file)
        print(f"Wrote {spectrogram_output_file}")
    return output_files


def make_comparison_plot(
    csv_files: list[Path] | None = None,
    observers: tuple[int, ...] = DEFAULT_OBSERVERS,
    plots: tuple[PlotDefinition, ...] = PLOTS,
    output_dir: Path = PLOTS_DIR,
    revolution_range: RevolutionRange | None = None,
    measurement_data: Poc2xMeasurementData | None = None,
    measurement_label: str = MEASUREMENT_LABEL,
) -> Path:
    csv_files = _default_acoustic_files() if csv_files is None else [Path(path) for path in csv_files]
    results_by_case = []
    for csv_file in csv_files:
        rpm, time_steps_per_revolution = _timing_for_file(csv_file, None, None)
        reference_density_kg_m3, reference_velocity_m_s, reference_temperature_k, l_grid_unit_m = _reference_values_for_file(
            csv_file,
            None,
            None,
            None,
            None,
        )
        dataset = load_acoustic_dataset(
            csv_file,
            observers,
            rpm=rpm,
            time_steps_per_revolution=time_steps_per_revolution,
            reference_density_kg_m3=reference_density_kg_m3,
            reference_velocity_m_s=reference_velocity_m_s,
            reference_temperature_k=reference_temperature_k,
            l_grid_unit_m=l_grid_unit_m,
            revolution_range=revolution_range,
        )
        values_by_plot = {plot.name: evaluate_plot(dataset, plot) for plot in plots}
        results_by_case.append((dataset, values_by_plot))
        print(
            f"Loaded {dataset.source_file.name}: "
            f"{len(dataset.time)} samples, "
            f"revolutions={dataset.revolution_range.label if dataset.revolution_range is not None else 'all'}, "
            f"rpm={dataset.rpm:.3f}, "
            f"steps/rev={dataset.time_steps_per_revolution:g}, "
            f"sample_rate={dataset.sample_rate_hz:.3f} Hz, "
            f"a={dataset.speed_of_sound_m_s:.6f} m/s, "
            f"L_grid_unit={dataset.l_grid_unit_m:g} m, "
            f"pressure_scale=gamma*p_inf={dataset.pressure_scale_pa:.3f} Pa"
        )
        if all(plot.name in values_by_plot for plot in PLOTS):
            print_oaspl_values(dataset, values_by_plot, observers)
        if (measurement_data is not None and "oaspl_a_weighted" in values_by_plot
                and "poc2" in {part.lower() for part in dataset.source_file.resolve().parts}):
            print_oaspl_differences(
                dataset,
                values_by_plot["oaspl_a_weighted"],
                measurement_data.oaspl_dba_by_observer,
                observers,
            )
            write_poc2_test_results(
                dataset,
                values_by_plot["oaspl_a_weighted"],
                measurement_data.oaspl_dba_by_observer,
                output_dir,
            )

    output_file = render_comparison_plot(
        results_by_case,
        observers,
        output_dir,
        measurement_oaspl_dba_by_observer=(
            measurement_data.oaspl_dba_by_observer if measurement_data is not None else None
        ),
        measurement_label=measurement_label,
    )
    print(f"Wrote {output_file}")
    for spectrum_output_file in render_spectrum_plots(results_by_case, observers, output_dir):
        print(f"Wrote {spectrum_output_file}")
    for third_octave_output_file in render_third_octave_plots(
        results_by_case,
        observers,
        output_dir,
        measurement_spectra_by_observer=(
            measurement_data.spectra_by_observer
            if measurement_data is not None and APPLY_A_WEIGHTING_TO_SPECTRA
            else None
        ),
        measurement_label=measurement_label,
    ):
        print(f"Wrote {third_octave_output_file}")
    for surface_pair_third_octave_output_file in render_surface_pair_third_octave_plots(
        results_by_case,
        observers,
        output_dir,
    ):
        print(f"Wrote {surface_pair_third_octave_output_file}")
    surface_pair_oaspl_output_file = render_surface_pair_oaspl_plots(
        results_by_case,
        observers,
        output_dir,
    )
    if surface_pair_oaspl_output_file is not None:
        print(f"Wrote {surface_pair_oaspl_output_file}")
    for bpf_output_file in render_bpf_component_plots(results_by_case, observers, output_dir):
        print(f"Wrote {bpf_output_file}")
    for harmonic_output_file in render_harmonic_content_plots(results_by_case, observers, output_dir):
        print(f"Wrote {harmonic_output_file}")
    for pressure_time_output_file in render_pressure_time_history_plots(results_by_case, observers, output_dir):
        print(f"Wrote {pressure_time_output_file}")
    for surface_pressure_time_output_file in render_surface_pressure_time_history_plots(
        results_by_case,
        observers,
        output_dir,
    ):
        print(f"Wrote {surface_pressure_time_output_file}")
    for average_pressure_time_output_file in render_average_pressure_time_history_plots(
        results_by_case,
        observers,
        output_dir,
    ):
        print(f"Wrote {average_pressure_time_output_file}")
    for spectrogram_output_file in render_spectrogram_plots(results_by_case, observers, output_dir):
        print(f"Wrote {spectrogram_output_file}")
    return output_file


def _parse_observers(value: str) -> tuple[int, ...]:
    observers = tuple(int(item.strip()) for item in value.split(",") if item.strip())
    if not observers:
        raise argparse.ArgumentTypeError("Provide at least one observer index.")
    return observers


def _parse_legend_labels(value: str) -> tuple[str, ...]:
    labels = tuple(item.strip() for item in value.split(",") if item.strip())
    if not labels:
        raise argparse.ArgumentTypeError("Provide at least one legend label.")
    return labels


def _parse_revolution_range(value: str) -> RevolutionRange | None:
    value = value.strip()
    if value.lower() == "all":
        return None

    match = re.fullmatch(r"(?:(zero|one|0|1):)?(\d+)\s*-\s*(\d+)", value, flags=re.IGNORECASE)
    if match is None:
        raise argparse.ArgumentTypeError(
            "Use a half-open revolution range such as 0-2 or 1-3. "
            "Use zero:1-3 or one:2-4 to force the index base."
        )

    base_token, start_text, end_text = match.groups()
    start = int(start_text)
    end = int(end_text)
    if base_token is None:
        index_base = 0 if start == 0 else 1
    elif base_token.lower() in {"zero", "0"}:
        index_base = 0
    else:
        index_base = 1

    normalized_start = start - index_base
    normalized_end = end - index_base
    if normalized_start < 0 or normalized_end <= normalized_start:
        raise argparse.ArgumentTypeError(
            f"Invalid revolution range {value!r}; the end must be greater than the start."
        )

    label = f"{value} ({normalized_start}-{normalized_end} zero-based, end-exclusive)"
    return RevolutionRange(
        start_index=normalized_start,
        end_index=normalized_end,
        label=label,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Plot Flow360 aeroacoustic OASPL results_postprocessing.")
    parser.add_argument(
        "--csv",
        type=Path,
        action="append",
        default=None,
        help="Path to *_results_total_acoustics_v*.csv. Can be passed more than once.",
    )
    parser.add_argument("--output-dir", type=Path, default=PLOTS_DIR)
    parser.add_argument("--combine-csv", action="store_true", help="Sum --csv files aligned by acoustic timestamps.")
    parser.add_argument("--n-timesteps-skip", type=int, default=0, help="Initial signal steps to remove per observer.")
    parser.add_argument("--skip-unit", choices=("cfd", "acoustic"), default="cfd")
    parser.add_argument("--observers", type=_parse_observers, default=DEFAULT_OBSERVERS)
    parser.add_argument(
        "--revolutions",
        type=_parse_revolution_range,
        default=None,
        help=(
            "Optional half-open revolution range used for all aeroacoustic calculations. "
            "Default uses all samples. Examples: 0-2 or 1-3 select the first two revolutions; "
            "zero:1-3 forces zero-based indexing."
        ),
    )
    parser.add_argument("--rpm", type=float, default=None)
    parser.add_argument("--steps-per-rev", type=float, default=None)
    parser.add_argument("--rho", type=float, default=None, help="Density for dimensionalizing acoustic pressure.")
    parser.add_argument(
        "--temperature-k",
        type=float,
        default=None,
        help="Static temperature for dimensionalizing acoustic pressure with gamma*p_inf.",
    )
    parser.add_argument("--u-ref", type=float, default=None, help="Reference velocity retained for reporting.")
    parser.add_argument(
        "--l-grid-unit",
        type=float,
        default=None,
        help="Flow360 grid-unit length in meters for dimensionalizing acoustic CSV time.",
    )
    parser.add_argument(
        "--csv-observer-offset",
        type=int,
        default=FLOW360_OBSERVER_NUMBERING_OFFSET,
        help="Offset from user observer number to CSV observer index. Default maps observer 1 to observer_0.",
    )
    parser.add_argument(
        "--legend-label",
        type=str,
        default=None,
        help="Override the plotted case label in legends.",
    )
    parser.add_argument(
        "--legend-labels",
        type=_parse_legend_labels,
        default=None,
        help="Comma-separated legend labels mapped to --csv order.",
    )
    parser.add_argument(
        "--measurement-excel",
        type=Path,
        default=None,
        help=(
            "Excel workbook containing POC2x measurement data. "
            "When provided, overlays 1000 RPM columns E:G on A-weighted 1/3-octave plots "
            "and row 28 OASPL on A-weighted observer plots for observers 2, 3, and 4. "
            "Also prints simulated-minus-measured A-weighted OASPL differences."
        ),
    )
    parser.add_argument(
        "--measurement-label",
        default=MEASUREMENT_LABEL,
        help="Legend label for the measured POC2 OASPL and third-octave spectra.",
    )
    weighting_group = parser.add_mutually_exclusive_group()
    weighting_group.add_argument(
        "--a-weighting",
        dest="apply_a_weighting_to_spectra",
        action="store_true",
        default=APPLY_A_WEIGHTING_TO_SPECTRA,
        help="Use A-weighting for spectrum and third-octave plots.",
    )
    weighting_group.add_argument(
        "--no-a-weighting",
        dest="apply_a_weighting_to_spectra",
        action="store_false",
        help="Use unweighted spectrum and third-octave plots.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.n_timesteps_skip < 0:
        raise ValueError("--n-timesteps-skip must be nonnegative.")
    if args.combine_csv and not args.csv:
        raise ValueError("--combine-csv requires explicit --csv files.")
    if args.n_timesteps_skip and args.csv and len(args.csv) > 1 and not args.combine_csv:
        raise ValueError("Multiple files with an initial cut require --combine-csv.")
    if args.combine_csv and args.legend_labels is not None:
        raise ValueError("Use --legend-label for the combined dataset, rather than --legend-labels.")
    global APPLY_A_WEIGHTING_TO_SPECTRA, FLOW360_OBSERVER_NUMBERING_OFFSET, LEGEND_LABEL_OVERRIDE, LEGEND_LABELS_BY_FILE
    APPLY_A_WEIGHTING_TO_SPECTRA = args.apply_a_weighting_to_spectra
    FLOW360_OBSERVER_NUMBERING_OFFSET = args.csv_observer_offset
    LEGEND_LABEL_OVERRIDE = args.legend_label
    if args.legend_labels is not None:
        if args.csv is None:
            raise ValueError("--legend-labels requires explicit --csv arguments.")
        if len(args.legend_labels) != len(args.csv):
            raise ValueError("--legend-labels must contain the same number of labels as --csv arguments.")
        LEGEND_LABELS_BY_FILE = {
            csv_file.resolve(): label
            for csv_file, label in zip(args.csv, args.legend_labels)
        }
    measurement_data = (
        load_poc2x_1000rpm_measurement_data(args.measurement_excel)
        if args.measurement_excel is not None
        else None
    )
    if (
        args.combine_csv
        or args.n_timesteps_skip
        or args.rpm is not None
        or args.steps_per_rev is not None
        or args.rho is not None
        or args.temperature_k is not None
        or args.u_ref is not None
        or args.l_grid_unit is not None
    ):
        if args.csv is not None and len(args.csv) > 1 and not args.combine_csv:
            raise ValueError(
                "Manual --rpm/--steps-per-rev/--rho/--temperature-k/--u-ref/--l-grid-unit overrides "
                "are only supported for one CSV."
            )
        make_plots(
            csv_file=args.csv[0] if args.csv else None,
            observers=args.observers,
            output_dir=args.output_dir,
            rpm=args.rpm,
            time_steps_per_revolution=args.steps_per_rev,
            reference_density_kg_m3=args.rho,
            reference_velocity_m_s=args.u_ref,
            reference_temperature_k=args.temperature_k,
            l_grid_unit_m=args.l_grid_unit,
            revolution_range=args.revolutions,
            measurement_data=measurement_data,
            measurement_label=args.measurement_label,
            csv_files=args.csv if args.combine_csv else None,
            n_timesteps_skip=args.n_timesteps_skip,
            skip_unit=args.skip_unit,
        )
    else:
        make_comparison_plot(
            csv_files=args.csv,
            observers=args.observers,
            output_dir=args.output_dir,
            revolution_range=args.revolutions,
            measurement_data=measurement_data,
            measurement_label=args.measurement_label,
        )
    if args.combine_csv or args.n_timesteps_skip:
        command = subprocess.list2cmdline([sys.executable, str(Path(__file__).resolve()), *sys.argv[1:]])
        with (args.output_dir / "commands.txt").open("a", encoding="utf-8") as file:
            file.write(command + "\n")


if __name__ == "__main__":
    main()
